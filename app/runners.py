# -*- coding: utf-8 -*-
"""Ban chep Accounting-Web\\runners.py cho app desktop Accounting Tools. Chi doi cho tim engine:
engine CHEP vao app\\engines\\{debt\\Source, pias\\tool, offset\\tool}, giu nguyen bo cuc workspace goc
nen config.xlsx / Templates\\ / Unit_Codes.xlsx / Offset_Rates.xlsx tu tim thay canh engine.
Sua engine o workspace goc thi CHEP LAI sang app\\engines.

worker.py dat RUNS / new_run vao thu muc cua lan chay, BACKUPS vao <goc app>\\backups.
"""
import contextlib
import datetime as dt
import glob
import io
import os
import re
import shutil
import sys
import threading
import time
import unicodedata
import uuid
import zipfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ENGINES = os.path.join(HERE, "engines")
DEBT_ROOT = os.path.join(ENGINES, "debt")
PIAS_ROOT = os.path.join(ENGINES, "pias")
OFFSET_ROOT = os.path.join(ENGINES, "offset")
HOME = os.path.dirname(HERE)
RUNS = os.path.join(HOME, "output")
BACKUPS = os.path.join(HOME, "backups")
KEEP_DAYS = 14

for _p in (os.path.join(DEBT_ROOT, "Source"), os.path.join(PIAS_ROOT, "tool"), os.path.join(OFFSET_ROOT, "tool")):
    if _p not in sys.path:
        sys.path.insert(0, _p)
os.environ["PIASFIX_NO_OPEN"] = "1"          # fix_pias khong tu mo Explorer
os.environ.pop("OFFSET_ROOT", None)          # offset_5111 tim Offset_Rates.xlsx canh tool\

import build_debt_report as B                # noqa: E402
import fix_pias as F                         # noqa: E402
try:
    import offset_5111 as O                  # noqa: E402
except ImportError:                          # goi local khong kem engine doi tru 5111
    O = None

ANSI = re.compile(r"\x1b\[[0-9;]*m")
XLSX_MAGIC = b"PK"


class RunError(Exception):
    """Loi do dau vao, hien thang cho nguoi dung (tieng Anh, trang tu dich)."""


# ------------------------------------------------------------------ hang doi

class JobQueue:
    """Hang doi FIFO dung chung CA APP (ca 2 trang): moi luc may chu chi chay mot viec.

    Anh Nam 2026-09-17: dang co nguoi chay bao cao cong no thi nguoi khac vao hang doi,
    khong duoc chay song song cho do lag may. Cung nho vay ma engine giu bien toan cuc
    (B.NHOM/GROUPS/GOP, F.OUT_ROOT) khong bi hai lan chay de len nhau.

    Ve rut khoi hang (finally) ca khi luot chay bi ngat giua chung: nguoi dung bam nut khac
    hay dong tab thi streamlit nem StopException/RerunException o lan goi st.* ke tiep.
    """

    def __init__(self):
        self._lock = threading.Lock()
        self._waiting = []
        self.running = None

    def join(self, label):
        t = {"id": uuid.uuid4().hex, "label": label, "joined": time.time(), "started": None}
        with self._lock:
            self._waiting.append(t)
        return t

    def try_start(self, t):
        with self._lock:
            if self.running is None and self._waiting and self._waiting[0] is t:
                self._waiting.pop(0)
                t["started"] = time.time()
                self.running = t
                return True
            return False

    def leave(self, t):
        with self._lock:
            if t in self._waiting:
                self._waiting.remove(t)
            if self.running is t:
                self.running = None

    def ahead(self, t):
        """So viec phai xong truoc t (tinh ca viec dang chay)."""
        with self._lock:
            i = self._waiting.index(t) if t in self._waiting else 0
            return i + (1 if self.running is not None else 0)

    def snapshot(self):
        """(viec dang chay hoac None, so nguoi dang cho) de hien o thanh ben."""
        with self._lock:
            return (dict(self.running) if self.running else None), len(self._waiting)


QUEUE = JobQueue()


@contextlib.contextmanager
def queued(label, on_wait=None, on_start=None, poll=1.0):
    """Xep hang roi chay. on_wait(so viec phia truoc, viec dang chay) goi moi giay khi con cho."""
    t = QUEUE.join(label)
    try:
        while not QUEUE.try_start(t):
            if on_wait:
                running, _ = QUEUE.snapshot()
                on_wait(QUEUE.ahead(t), running)
            time.sleep(poll)
        if on_start:
            on_start()
        yield t
    finally:
        QUEUE.leave(t)


# ------------------------------------------------------------------ tien ich chung

def nfc(s):
    return unicodedata.normalize("NFC", str(s))


def new_run(kind):
    base = os.path.join(RUNS, kind)
    os.makedirs(base, exist_ok=True)
    cutoff = time.time() - KEEP_DAYS * 86400
    for n in os.listdir(base):
        p = os.path.join(base, n)
        try:
            if os.path.isdir(p) and os.path.getmtime(p) < cutoff:
                shutil.rmtree(p, ignore_errors=True)
        except OSError:
            pass
    d = os.path.join(base, dt.datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:6])
    os.makedirs(d)
    return d


def save_files(files, folder, rename=None):
    """files = [(ten, bytes)]. rename(ten) -> ten moi hoac None. Trung ten thi them (2)."""
    os.makedirs(folder, exist_ok=True)
    out = []
    for name, data in files:
        name = nfc(os.path.basename(name))
        name = (rename(name) if rename else None) or name
        stem, ext = os.path.splitext(name)
        p, i = os.path.join(folder, name), 2
        while os.path.exists(p):
            p, i = os.path.join(folder, "%s (%d)%s" % (stem, i, ext)), i + 1
        with open(p, "wb") as f:
            f.write(data)
        out.append(p)
    return out


def zip_bytes(paths):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
    return buf.getvalue()


def outputs_in(folder):
    if not os.path.isdir(folder):
        return []
    return sorted(p for p in glob.glob(os.path.join(folder, "*"))
                  if os.path.isfile(p) and not os.path.basename(p).startswith(("~$", "_recalc_")))


def replace_file(target, data, check):
    """Thay file cai dat (config.xlsx, Unit_Codes.xlsx) sau khi check(data) qua; ban cu vao backups\\."""
    if data[:2] != XLSX_MAGIC:
        raise RunError("This is not an .xlsx file.")
    check(data)
    os.makedirs(BACKUPS, exist_ok=True)
    stem, ext = os.path.splitext(os.path.basename(target))
    if os.path.exists(target):
        shutil.copy2(target, os.path.join(BACKUPS, "%s_%s%s" % (stem, dt.datetime.now().strftime("%Y%m%d_%H%M%S"), ext)))
    tmp = target + ".upload.tmp"
    with open(tmp, "wb") as f:
        f.write(data)
    try:
        os.replace(tmp, target)
    except PermissionError:
        os.remove(tmp)
        raise RunError("The file is open on the server, try again in a minute.")


# ------------------------------------------------------------------ BC cong no 1311

DEBT_KINDS = ("raw", "bk", "list", "report")
LIST_NAME = "List hóa đơn.xlsx"      # dung ten nay de engine nhan ra (invoices.MAU_LIST)
# chay le (ke toan 30/09/2026: check rieng 0051 theo phong KD, khoi chay ca 13111)
DEBT_SCOPES = {"full": B.PHAN_DU, "0051": ("0051",), "13111": ("13111",)}
SCOPE_ACCOUNTS = {"full": tuple(B.ACCOUNTS), "0051": ("0051",),
                  "13111": ("131112", "131113", "131114", "131115")}


def _plain(s):
    s = unicodedata.normalize("NFD", nfc(s)).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


def debt_kind(name):
    k = _plain(os.path.basename(name))
    if k.startswith("bcpbhtt"):
        return "raw"
    if k.startswith("bk_hdvat"):
        return "bk"
    if k.startswith("list hoa don"):
        return "list"
    if k.startswith("bc cong no"):
        return "report"
    return None


def raw_account(data):
    """Ma tai khoan cua 1 file BCPBHTT, doc o dong 6 nhu engine (B.read_raw). None neu khong ra."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        for row in wb.worksheets[0].iter_rows(min_row=6, max_row=6, max_col=8, values_only=True):
            for v in row:
                m = re.search(r":\s*([0-9]{4,6})", str(v or ""))
                if m:
                    return m.group(1)
        return None
    finally:
        wb.close()


def missing_accounts(scope, accounts):
    """Tai khoan pham vi nay can ma chua co file. accounts = ma TK doc duoc tu cac file BCPBHTT."""
    return [a for a in SCOPE_ACCOUNTS[scope] if a not in accounts]


def debt_problems(counts, clean, kh, pkd, scope="full"):
    """Vi sao chua chay duoc (danh sach cau tieng Anh), rong = chay duoc.
    Thieu file tai khoan nao thi trang tu bao (missing_accounts), engine cung chan lai."""
    out = []
    if not counts.get("raw") and not counts.get("report"):
        out.append("Upload the BCPBHTT files, or a BC công nợ file that is already built.")
    if kh and scope == "0051":
        out.append("Letters to customers use account 13111, so they cannot run with 0051 only.")
    if clean and not (counts.get("bk") and counts.get("list")):
        out.append("Clean invoices needs both BK_HDVAT and List hóa đơn.")
    if kh and not counts.get("list"):
        out.append("Letters to customers needs List hóa đơn.")
    for k in ("bk", "list", "report"):
        if counts.get(k, 0) > 1:
            out.append({"bk": "Upload only one BK_HDVAT file.",
                        "list": "Upload only one List hóa đơn file.",
                        "report": "Upload only one BC công nợ file."}[k])
    return out


def debt_run(files, verify=True, clean=False, kh=False, pkd=False, log=print, progress=None,
             on_wait=None, on_start=None, scope="full"):
    """files = [(ten, bytes)]. Tra ve dict: run, outputs, report, verified (None/True/False), log_text.
    scope: "full" (13111 + 0051), "0051" hay "13111" de chay le; file TK ngoai pham vi bi bo qua.
    Chay qua hang doi chung (queued): on_wait/on_start de trang hien vi tri trong hang."""
    kinds = [debt_kind(n) for n, _ in files]
    counts = {k: kinds.count(k) for k in DEBT_KINDS}
    bad = debt_problems(counts, clean, kh, pkd, scope)
    if bad:
        raise RunError(" ".join(bad))
    lines = []

    def say(m, tag=None):
        lines.append(str(m))
        log(m)

    with queued("Debt report 1311", on_wait, on_start):
        run = new_run("debt")
        inp, out = os.path.join(run, B.DIR_IN), os.path.join(run, B.DIR_OUT)
        os.makedirs(out)
        paths = save_files(files, inp, rename=lambda n: LIST_NAME if debt_kind(n) == "list" else None)
        by = {k: [p for p in paths if debt_kind(p) == k] for k in DEBT_KINDS}
        bk = by["bk"][0] if by["bk"] else None
        lst = by["list"][0] if by["list"] else None
        verified = None
        try:
            if by["raw"]:
                rp = B.build(inp, out, log=say, progress=progress, phan=DEBT_SCOPES[scope])
                if verify:
                    import verify_report as V
                    if progress:
                        progress(96, "Excel đang tính lại và đối chiếu")
                    com = _com_init()
                    try:
                        verified = bool(V.verify(rp, inp, log=say))
                    finally:
                        _com_done(com)
            else:
                rp = by["report"][0]
                say("Dung bao cao da tai len: %s" % os.path.basename(rp))
            if clean:
                import build_invoice_master as M
                M.dung(run, log=say, f_bk=bk, f_li=lst, progress=progress)
            lam = tuple(k for k, on in (("kh", kh), ("pkd", pkd)) if on)
            if lam:
                import build_letters as LT
                LT.build_all(rp, out, inp, log=say, progress=progress, lam=lam, inv_path=lst, bk_path=bk)
        except SystemExit as e:                  # engine bao loi dau vao bang SystemExit(message)
            raise RunError(str(e)) from None
        finally:
            with io.open(os.path.join(out, "run-log.txt"), "w", encoding="utf-8") as f:
                f.write("\n".join(lines) + "\n")
        if progress:
            progress(100, "Xong")
    return {"run": run, "outputs": outputs_in(out), "report": rp,
            "verified": verified, "log_text": "\n".join(lines)}


def _com_init():
    try:
        import pythoncom
        pythoncom.CoInitialize()               # luong cua streamlit chua khoi tao COM cho Excel
        return pythoncom
    except ImportError:
        return None


def _com_done(pythoncom):
    if pythoncom:
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def debt_config_path():
    return str(B.config_path())


def replace_debt_config(data):
    need = {"Nhóm nghiệp vụ", "Gộp khách hàng", "Phòng KD", "Tham số", "Thông tin KH"}

    def check(d):
        try:
            names = {nfc(s) for s in openpyxl.load_workbook(io.BytesIO(d), read_only=True).sheetnames}
        except Exception:
            raise RunError("Cannot open this file as a workbook.")
        miss = sorted(need - names)
        if miss:
            raise RunError("This does not look like config.xlsx, missing sheet(s): %s." % ", ".join(miss))
    replace_file(debt_config_path(), data, check)


# ------------------------------------------------------------------ sua file thu PIAS

def pias_kind(data):
    return "pias" if F.is_pias(io.BytesIO(data)) else "list"


SUMMARY_RE = {
    "file": re.compile(r"^\[(?P<f>.+?)\]\s+(?P<rows>\d+) rows\s+\|\s+(?P<fixed>\d+) fixed\s+\|\s+(?P<done>\d+) already fixed", re.M),
    "check": re.compile(r"!! (\d+) to check by hand"),
    "look": re.compile(r"!! (\d+) fixed, but please look"),
    "left": re.compile(r"!! (\d+) item\(s\) of the PVI khac list NOT matched"),
}


def pias_run(files, sheet=None, dates=None, on_wait=None, on_start=None):
    """files = [(ten, bytes)]. Chay y nhu Run_PIAS_Fix.bat tren mot thu muc input tam.
    Tra ve dict: run, error, console, outputs, log_text, files (tung file PIAS), check, look, left."""
    kinds = [pias_kind(d) for _, d in files]
    if "pias" not in kinds:
        raise RunError("No PIAS file among the uploads (the columns 'Diễn giải' and 'Tk có' were not found).")
    if "list" not in kinds:
        raise RunError("Upload the PVI khác list from Ms. Kiều as well.")
    args = []
    if sheet:
        args += ["--sheet", sheet.strip()]
    if dates:
        args += ["--date", dates.strip()]
    with queued("PIAS fix for other PVI units", on_wait, on_start):
        run = new_run("pias")
        inp = os.path.join(run, "input")
        save_files(files, inp)
        buf, err = io.StringIO(), None
        old = F.OUT_ROOT, F.LAST_DIR
        F.OUT_ROOT, F.LAST_DIR = os.path.join(run, "output"), os.path.join(run, "last_folder.txt")
        try:
            with contextlib.redirect_stdout(buf):
                F.main([inp] + args)
        except F.ToolError as e:
            err = str(e)
        except SystemExit:
            err = "Cannot read the options. Sheet is a sheet name; days look like 08/09,09/09."
        finally:
            F.OUT_ROOT, F.LAST_DIR = old
    console = _web_console(ANSI.sub("", buf.getvalue()))
    outs =[p for d in glob.glob(os.path.join(run, "output", "*")) for p in outputs_in(d)]
    log_path = next((p for p in outs if os.path.basename(p) == "log.txt"), None)
    log_text = io.open(log_path, encoding="utf-8").read() if log_path else ""
    return {
        "run": run, "error": err, "console": console.strip(), "log_text": log_text,
        "outputs": sorted(outs, key=lambda p: (os.path.basename(p) == "log.txt", os.path.basename(p))),
        "files": [m.groupdict() for m in SUMMARY_RE["file"].finditer(console)],
        "check": sum(int(x) for x in SUMMARY_RE["check"].findall(console)),
        "look": sum(int(x) for x in SUMMARY_RE["look"].findall(console)),
        "left": sum(int(x) for x in SUMMARY_RE["left"].findall(console)),
    }


def _web_console(text):
    """Bo dong duong dan tren may chu (Folder / Output folder): tren web file tai ve bang nut."""
    out, skip = [], False
    for line in text.splitlines():
        if skip:
            skip = False
            continue
        if line.startswith("Folder   :"):
            continue
        if line.startswith("Output folder:"):
            skip = True
            continue
        out.append(line)
    return "\n".join(out)


def unit_codes_path():
    return F.UNIT_FILE


def replace_unit_codes(data):
    def check(d):
        try:
            ws = openpyxl.load_workbook(io.BytesIO(d), read_only=True).worksheets[0]
            rows = [r for r in ws.iter_rows(min_row=2, max_col=2, values_only=True) if r and r[0] and r[1]]
        except Exception:
            raise RunError("Cannot open this file as a workbook.")
        good = [r for r in rows if F.CODE_RE.match(str(r[1]).strip())]
        if len(good) < 5 or len(good) < len(rows) * 0.8:
            raise RunError("Expected 2 columns: unit name, then a code like 00.18520000, from row 2 down.")
    replace_file(unit_codes_path(), data, check)


# ------------------------------------------------------------------ doi tru voi so 5111

def has_offset():
    """Co engine doi tru 5111 hay khong (goi local khong kem engine nay)."""
    return O is not None


def offset_run(ledger, detail, chot=None, ky=None, on_wait=None, on_start=None, progress=None):
    """ledger, detail = (ten, bytes). Ap boc tach roi doi tru, tra ve dict cho trang web.

    Anh Nam 2026-10-14: trang co HAI o tai file rieng, KHONG do loai file nua (mo header cua so
    5111 nang 26 MB ton vai giay moi lan ve lai trang). File dua sai cho thi engine bao ngay:
    read_ledger / read_detail nem ToolError noi ro thieu cot gi.

    chot: ngay chot ky cho cot mvv (None = cuoi thang cua Ngay ctu moi nhat trong so).
    ky  : chu "Ky thanh toan" in tren sheet BBDS (None = lay theo Ngay cap don cua file chi tiet).
    """
    if O is None:
        raise RunError("The offset engine is not installed in this copy.")
    if not ledger or not detail:
        raise RunError("Upload the 5111 ledger and the detail file.")
    lines = []

    def say(m):
        lines.append(str(m))

    with queued("Offset 5111", on_wait, on_start):
        run = new_run("offset")
        inp, out = os.path.join(run, "input"), os.path.join(run, "output")
        p_led, p_det = save_files([ledger], inp)[0], save_files([detail], inp)[0]
        try:
            res = O.run(p_led, p_det, out, chot=chot, ky=ky, log=say, progress=progress)
        except O.ToolError as e:
            raise RunError(str(e)) from None
        finally:
            with io.open(os.path.join(run, "run-log.txt"), "w", encoding="utf-8") as f:
                f.write("\n".join(lines) + "\n")
    res["run"] = run
    res["outputs"] = outputs_in(out)
    res["log_text"] = "\n".join(lines)
    return res


def offset_rates_path():
    if O is None:
        raise RunError("The offset engine is not installed in this copy.")
    return O.rates_path()


def replace_offset_rates(data):
    def check(d):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(d), read_only=True, data_only=True)
        except Exception:
            raise RunError("Cannot open this file as a workbook.")
        try:
            names = {"ma_sp": ["Mã sp"], "hh": ["Tỷ lệ hoa hồng"], "pdv": ["Tỷ lệ phí dịch vụ"]}
            found = O.find_header(wb.worksheets[0], names, ("ma_sp", "hh", "pdv"))
            if not found:
                raise RunError("Expected the columns Mã sp, Tỷ lệ hoa hồng and Tỷ lệ phí dịch vụ "
                               "on the first sheet.")
            head_row, _header, cols = found
            ws = wb.worksheets[0]
            n = 0
            for row in ws.iter_rows(min_row=head_row + 1, values_only=True):
                sp = row[cols["ma_sp"]] if cols["ma_sp"] < len(row) else None
                if re.fullmatch(r"\d{4,8}", O.norm_sp(sp)):
                    n += 1
            if n < 1:
                raise RunError("No product row found under the header.")
        finally:
            wb.close()
    replace_file(offset_rates_path(), data, check)
