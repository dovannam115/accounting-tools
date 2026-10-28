"""Xuat 2 file tu bao cao cong no: thu gui khach hang va ban doi chieu gui P.KD.

    py -3.13 build_letters.py <file BC cong no> [thu muc ghi ket qua]

Giu nguyen dinh dang 2 file mau bang cach nhan ban chinh sheet cua chung.
"""
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from math import ceil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.cell_range import CellRange

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_debt_report import (DIR_IN, DIR_TPL, app_dir, clean_name,  # noqa: E402
                               config_path, file_kem, norm, vn_key)
from build_pkd_annex import build_annex, loc_0051                             # noqa: E402
from invoices import chon_hoa_don, doc_hoa_don, tim_file                      # noqa: E402
from vn_number import doc_tien                                                # noqa: E402

TPL_KH = "gửi thư KH.xlsx"
TPL_PKD = "BẢN ĐỐI CHIẾU CÔNG NỢ GỬI P.KD T06.2025.xlsx"

# cot trong sheet bao cao (0-based) — 24 cot dau giong nhau o BC 13111 va 0051
C_MA_KH, C_TEN_KH, C_SO_DON, C_PHONG = 1, 2, 4, 7
C_DAI_LY, C_NGAY_CAP, C_PHAI_THU, C_DA_THU, C_CON_LAI = 10, 13, 20, 21, 22
C_LOAI_TIEN = 15
ROW_DATA = 12

# vi tri trong template thu KH
KH_ROW_DATA, KH_ROW_END = 24, 42
# vi tri trong template P.KD
PKD_ROW_DATA, PKD_ROW_END = 18, 41
PKD_TPL_M1, PKD_TPL_M2 = 2, 2          # so dong 13111 / 0051 co san trong mau


def q_cua(d):
    return f"quý {(d.month - 1) // 3 + 1:02d}/{d.year}"


def doc_sheet(path, ten_sheet):
    """Cac dong du lieu cua sheet; None khi file khong co sheet do (bao cao chay le)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if ten_sheet not in wb.sheetnames:
        wb.close()
        return None
    ws = wb[ten_sheet]
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True), 1):
        if i >= ROW_DATA and r and len(r) > C_CON_LAI and r[C_MA_KH] not in (None, ""):
            out.append(r)
    wb.close()
    return out


def doc_config():
    wb = load_workbook(config_path(), data_only=True)
    kh_ma, kh_ten = {}, {}
    if "Thông tin KH" in wb.sheetnames:
        for r in wb["Thông tin KH"].iter_rows(min_row=3, values_only=True):
            ma, ten, mst, dc = (tuple(r) + (None,) * 4)[:4]
            if not ma:
                continue
            rec = (str(mst or "").strip(), str(dc or "").strip())
            kh_ma[str(ma).strip()] = rec
            if ten:
                kh_ten.setdefault(vn_key(ten), rec)
    ts = dict(nguong=0.0, chi_qua_han=False, bo_khong=True)
    if "Tham số" in wb.sheetnames:
        for r in wb["Tham số"].iter_rows(min_row=3, values_only=True):
            ten, gt = (tuple(r) + (None, None))[:2]
            t = norm(ten)
            co = str(gt or "").strip().lower().startswith("c")
            if t.startswith("DƯ NỢ TỐI THIỂU"):
                ts["nguong"] = abs(float(gt or 0))
            elif t.startswith("CHỈ LẬP THƯ"):
                ts["chi_qua_han"] = co
            elif t.startswith("BỎ THƯ"):
                ts["bo_khong"] = co
    phong = {}
    if "Phòng KD" in wb.sheetnames:
        for r in wb["Phòng KD"].iter_rows(min_row=3, values_only=True):
            pb, sheet, ten, xh, nguoi, cv, nhan, gom = (tuple(r) + (None,) * 8)[:8]
            if not pb:
                continue
            phong[norm(pb)] = dict(pias=str(pb).strip(), sheet=str(sheet or "").strip(),
                                   ten=str(ten or "").strip(), xh=str(xh or "Ông").strip(),
                                   nguoi=str(nguoi or "").strip(), cv=str(cv or "").strip(),
                                   nhan=str(nhan or "").strip(),
                                   gom=str(gom or "").strip().lower())
    wb.close()
    return kh_ma, kh_ten, phong, ts


def nguoi_nhan(r):
    """Khach le co dai ly -> gom ve dai ly. Con lai (ke ca chi nhanh) -> gui rieng."""
    ten = clean_name(r[C_TEN_KH])
    dl = clean_name(r[C_DAI_LY])
    if ten.lower().startswith("khách lẻ") and dl:
        return dl, None, True
    return ten, str(r[C_MA_KH] or "").strip(), False


def ten_sheet_hop_le(ten, da_dung):
    s = re.sub(r"[\[\]:*?/\\]", " ", clean_name(ten))[:31].strip() or "KH"
    g, i = s, 2
    while g.lower() in da_dung:
        hau = f" ({i})"
        g = s[:31 - len(hau)] + hau
        i += 1
    da_dung.add(g.lower())
    return g


def _thay_ngay(v, ngay, nam_mau, ngay_mau, q_mau=None):
    if not isinstance(v, str):
        return v
    v = v.replace(ngay_mau, f"{ngay:%d/%m/%Y}").replace(f"năm {nam_mau}", f"năm {ngay.year}")
    if q_mau:
        v = v.replace(q_mau, q_cua(ngay))
    return v


def _don_khoi(ws, tu_dong):
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= tu_dong:
            ws.unmerge_cells(str(m))
    for r in range(tu_dong, ws.max_row + 2):
        for c in range(1, 12):
            cell = ws.cell(r, c)
            if not isinstance(cell, MergedCell):
                cell.value = None


def _cau_truc_pkd(ws):
    """7 sheet mau khong cung bo cuc dong — phai do tung sheet.
    -> (dong data 13111, dong Tong 1, dong data 0051, dong Tong 2, dong TONG CONG, dong Bang chu)"""
    r_head = next((r for r in range(1, 40)
                   if str(ws.cell(r, 1).value or "").strip() == "Tài khoản"), None)
    if not r_head:
        return None
    def tim_tong(tu):
        for r in range(tu, tu + 60):
            if "Tổng cộng" in str(ws.cell(r, 1).value or ""):
                return r
        return None
    d1 = r_head + 2
    t1 = tim_tong(d1)
    if not t1:
        return None
    d2 = t1 + 1
    t2 = tim_tong(d2)
    if not t2:
        return None
    return d1, t1, d2, t2, t2 + 1, t2 + 2


def _sheet_phong_moi(wb, sn, p, ten_sheet, log_rows):
    """Config khai phong moi ma file mau chua co sheet -> nhan ban mot sheet phong co san
    roi thay ten phong / nguoi dai dien / chuc vu / nhan ky bang so khai trong config.

    KHONG BAO GIO de nguyen ten nguoi cua sheet nguon: khai thieu thi de trong va ghi log,
    chu in nham ten nguoi ky vao van ban chinh thuc con te hon la de trong.
    """
    nguon = next((s for s in wb.sheetnames
                  if ten_sheet.get(s.strip()) and _cau_truc_pkd(wb[s])), None)
    if not nguon:
        return None
    goc = ten_sheet[nguon.strip()]
    ws = wb.copy_worksheet(wb[nguon])
    ws.title = sn
    ws.sheet_view.showGridLines = wb[nguon].sheet_view.showGridLines
    ws.print_title_rows = wb[nguon].print_title_rows

    thieu = [ten for ten, k in (("Tên phòng ghi trong thư", "ten"), ("Người đại diện", "nguoi"),
                                ("Chức vụ", "cv"), ("Nhãn ký", "nhan")) if not p.get(k)]
    thay = [(goc["cv"], p["cv"]), (goc["nhan"], p["nhan"] or f"ĐẠI DIỆN {sn}"),
            (goc["ten"], p["ten"] or p["pias"]), (goc["pias"], p["ten"] or p["pias"]),
            (goc["nguoi"], p["nguoi"])]
    thay = sorted([(a, b) for a, b in thay if a], key=lambda x: -len(x[0]))
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            v = c.value
            for a, b in thay:
                if a in v:
                    v = v.replace(a, b)
                    break
            if v != c.value:
                c.value = v
    # "Ong/Ba" phai thay theo VI TRI: ten cua P.Ke toan o tren cung dung chu 'Ông:',
    # replace theo chuoi se sua nham dong do.
    if goc["nguoi"]:
        for row in ws.iter_rows(min_col=2, max_col=2):
            if row[0].value == p["nguoi"] or (not p["nguoi"] and row[0].value == ""):
                o = ws.cell(row[0].row, 1)
                if isinstance(o.value, str) and o.value.strip().rstrip(":") in ("Ông", "Bà"):
                    o.value = p["xh"] + (":" if o.value.strip().endswith(":") else "")
                break
    log_rows.append(("Tự tạo sheet mới từ file mẫu, soát lại câu chữ trước khi gửi",
                     "", "", f"{sn} (nhân bản từ {nguon.strip()})", "P.KD",
                     p["ten"] or p["pias"], "", ""))
    if thieu:
        log_rows.append((f"Sheet mới thiếu {', '.join(thieu)} trong config.xlsx, ô để trống",
                         "", "", sn, "P.KD", p["ten"] or p["pias"], "", ""))
    return ws


def _ap_config_vao_sheet(ws, p, sn, log_rows):
    """Ap ten phong / nguoi ky khai trong config.xlsx vao MOT sheet ban xac nhan.

    Sheet co san trong file mau mang ten nguoi ky cua ky mau (T06.2025) va truoc day
    khong ai ghi de -> sua cot 'Nguoi dai dien' / 'Chuc vu' / 'Nhan ky' trong config
    ma phong da co sheet thi ban in ra van la ten nguoi cu. Chi sheet TU TAO moi
    (_sheet_phong_moi) mo doc config.

    Do theo VI TRI, khong replace chuoi: khoi 'Dai dien:' THU HAI moi la phong KD,
    khoi dau tien la P.Ke toan (sua nham la doi ten truong phong ke toan).
    Config bo trong o nao thi GIU nguyen o do cua file mau + ghi log — khac voi sheet
    tu tao (nhan ban cua phong khac nen bat buoc phai xoa ten nguoi).
    """
    dr = [r for r in range(1, min(ws.max_row, 30) + 1)
          if str(ws.cell(r, 1).value or "").strip().rstrip(":").lower() == "đại diện"]
    if len(dr) < 2:
        log_rows.append(("Không dò được khối chữ ký trong sheet mẫu, tên người ký giữ theo file mẫu",
                         "", "", sn, "P.KD", p["ten"] or p["pias"], "", ""))
        return
    r = dr[1]
    thieu = []

    def dat(o, gt, ten_cot, giu_duoi=""):
        if not gt:
            thieu.append(ten_cot)
            return
        o.value = gt + giu_duoi

    ten = p["ten"] or p["pias"]
    if str(ws.cell(r - 1, 1).value or "").strip():
        ws.cell(r - 1, 1).value = ten
    o_xh = ws.cell(r + 1, 1)
    cu = str(o_xh.value or "").strip()
    if cu.rstrip(":") in ("Ông", "Bà"):
        o_xh.value = p["xh"] + (":" if cu.endswith(":") else "")
    dat(ws.cell(r + 1, 2), p["nguoi"], "Người đại diện")
    o_cv = ws.cell(r + 1, 4)
    dau = "Chức vụ: " if str(o_cv.value or "").strip().startswith("Chức vụ") else ""
    if p["cv"]:
        o_cv.value = dau + p["cv"]
    else:
        thieu.append("Chức vụ")

    # Khoi ky ten cuoi trang: 2 cot canh nhau, cot TRAI la P.Ke toan, cot PHAI la phong KD.
    nhan = [c for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.strip().startswith("ĐẠI DIỆN")]
    if nhan:
        o = max(nhan, key=lambda c: (c.column, c.row))
        dat(o, p["nhan"], "Nhãn ký")
        for rr in range(o.row + 1, ws.max_row + 1):
            c = ws.cell(rr, o.column)
            if isinstance(c.value, str) and c.value.strip():
                dat(c, p["nguoi"], "Người đại diện")
                break
    if thieu:
        log_rows.append((f"config.xlsx chưa khai {', '.join(sorted(set(thieu)))}, ô giữ theo file mẫu",
                         "", "", sn, "P.KD", ten, "", ""))


def _sua_chieu_cao(ws, tu_dong, den_dong, ncol=7):
    """Excel khong tu gian dong cho o da merge — phai tu tinh chieu cao, khong thi
    chu dai bi cat mat luc in."""
    rong = {c: (ws.column_dimensions[L(c)].width or 8.43) for c in range(1, ncol + 1)}
    nhip = {}
    for m in ws.merged_cells.ranges:
        if m.min_row == m.max_row:
            nhip.setdefault(m.min_row, []).append((m.min_col, m.max_col))
    for r in range(tu_dong, den_dong + 1):
        can = 0
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(cell, MergedCell) or not isinstance(v, str) or not v.strip():
                continue
            if not cell.alignment.wrap_text:
                continue
            a, b = next((s for s in nhip.get(r, []) if s[0] <= c <= s[1]), (c, c))
            w = sum(rong.get(i, 8.43) for i in range(a, b + 1))
            sz = cell.font.sz or 11
            cpl = max(8, w * 10.8 / sz)          # so ky tu vua mot dong
            dong = sum(max(1, ceil(len(d) / cpl)) for d in v.split("\n"))
            if dong > 1:
                can = max(can, round(dong * sz * 1.35, 1))
        if can and can > (ws.row_dimensions[r].height or 0):
            ws.row_dimensions[r].height = can


def _chep_dong(ws, tpl, r_dich, r_nguon, ncol=10):
    for c in range(1, ncol + 1):
        ws.cell(r_dich, c)._style = tpl.cell(r_nguon, c)._style
    h = tpl.row_dimensions[r_nguon].height
    if h:
        ws.row_dimensions[r_dich].height = h


# ------------------------------------------------------------------ thu KH ---
def build_kh(bang, inv, kh_ma, kh_ten, ts, ngay, out_dir, log=print):
    tpl_path = file_kem(TPL_KH, DIR_TPL)
    if not tpl_path.exists():
        raise SystemExit(f"Khong thay file mau:\n{tpl_path}")
    wb = load_workbook(tpl_path)
    tpl = wb[wb.sheetnames[0]]
    merges = [str(m) for m in tpl.merged_cells.ranges]
    mau = {r: [tpl.cell(r, c).value for c in range(1, 11)] for r in range(25, KH_ROW_END + 1)}

    theo_kh = defaultdict(list)
    for sheet, rows in bang.items():
        for r in rows:
            con = r[C_CON_LAI] if isinstance(r[C_CON_LAI], (int, float)) else 0
            if abs(con) < 0.005:
                continue
            ten, ma, gom = nguoi_nhan(r)
            theo_kh[ten].append((sheet, r, ma, gom))

    def giu_lai(muc):
        tong = sum(r[C_CON_LAI] for _s, r, _m, _g in muc if isinstance(r[C_CON_LAI], (int, float)))
        qh = sum(r[23] for _s, r, _m, _g in muc if isinstance(r[23], (int, float)))
        if ts["bo_khong"] and abs(tong) < 0.005:
            return False, "tổng dư nợ bằng 0"
        if ts["chi_qua_han"] and abs(qh) < 1:
            return False, "không có nợ quá hạn"
        if abs(tong) < ts["nguong"]:
            return False, f"dư nợ dưới ngưỡng {ts['nguong']:,.0f}"
        return True, ""

    log_rows, da_dung = [], set()
    bo = 0
    for ten in sorted(theo_kh, key=vn_key):
        muc = theo_kh[ten]
        giu, ly_do = giu_lai(muc)
        if not giu:
            bo += 1
            log_rows.append((f"Không lập thư, {ly_do}", "", "", "", "thư gửi KH", ten,
                             sum(r[C_CON_LAI] for _s, r, _m, _g in muc
                                 if isinstance(r[C_CON_LAI], (int, float))), ""))
            continue
        # Mot so don co the nam nhieu dong (13111 + 0051, hoac lap trong cung 0051) —
        # gop lai 1 dong cho de doi chieu, dong nao cong xong ra 0 thi bo han.
        gd = {}
        for _sheet, r, _ma, _g in muc:
            k = str(r[C_SO_DON] or "").strip()
            d = gd.setdefault(k, dict(u=0.0, v=0.0, w=0.0, n=0, tte=set()))
            for phim, cot in (("u", C_PHAI_THU), ("v", C_DA_THU), ("w", C_CON_LAI)):
                if isinstance(r[cot], (int, float)):
                    d[phim] += r[cot]
            d["n"] += 1
            if str(r[C_LOAI_TIEN] or "").strip().upper() not in ("", "VND"):
                d["tte"].add(str(r[C_LOAI_TIEN]).strip().upper())
        dong = sorted(((k, d) for k, d in gd.items() if abs(d["w"]) >= 0.005),
                      key=lambda x: x[0])
        if not dong:
            bo += 1
            log_rows.append(("Không lập thư, mọi dòng cộng lại bằng 0", "", "", "",
                             "thư gửi KH", ten, "", ""))
            continue

        ten_sh = ten_sheet_hop_le(ten, da_dung)
        ws = wb.copy_worksheet(tpl)
        ws.title = ten_sh

        ma_kh = next((m for _s, _r, m, g in muc if m and not g), None)
        mst, dia_chi = kh_ma.get(ma_kh or "", ("", ""))
        if not mst and not dia_chi:
            mst, dia_chi = kh_ten.get(vn_key(ten), ("", ""))
        if not mst and not dia_chi:
            log_rows.append(("Không tra được MST/địa chỉ", "", "", ten_sh, "thư gửi KH", ten, "", ""))

        n = len(dong)
        _don_khoi(ws, KH_ROW_DATA)
        ws["B13"] = ten
        ws["B14"] = mst or None
        ws["B15"] = dia_chi or None
        ws["A19"] = _thay_ngay(ws["A19"].value, ngay, 2026, "30/06/2026", "quý 02/2026")

        tong_hd = tong_con = 0.0
        for i, (so_don, d) in enumerate(dong):
            rr = KH_ROW_DATA + i
            _chep_dong(ws, tpl, rr, KH_ROW_DATA)
            nhan, tien, _ng, so_ngay = chon_hoa_don(inv, so_don)
            phai_thu, da_thu, con = d["u"], d["v"], d["w"]
            ws.cell(rr, 1, i + 1)
            ws.cell(rr, 2, so_don)
            c3 = ws.cell(rr, 3, nhan)
            # mot don co the co nhieu so hoa don -> cho xuong dong de khong bi cat
            c3.alignment = Alignment(horizontal=c3.alignment.horizontal or "center",
                                     vertical="center", wrap_text=True)
            ws.cell(rr, 4, tien if nhan else None)
            # F = "So da thanh toan / So tien" = cot V cua PIAS (so_tien_tt). Dung duoc vi
            # 'Phai thu = Da thu + Con lai' khop tren 8.356/8.523 dong (167 dong lech vai chuc
            # dong do quy doi ngoai te), va thu chi lay TK 13111 nen khong con canh mot don
            # nam ca 13111 lan 0051 gay cong doi. Ky 31/07 chi 94 dong co so.
            # E "Ngay" DE TRONG: PIAS chi co 'Ngay thu phi' (ngay_ttoan) la HAN thu phi —
            # don cap 11/07 chua thu dong nao van ghi 16/08 — khong phai ngay khach tra.
            # DON NGOAI TE -> in con no theo HOA DON, khong theo so. So cong no da duoc danh
            # gia lai theo ty gia cuoi ky con khach ghi so dung so VND tren hoa don, vd
            # 25/45/06/CHBB/PC00206/E11 = 54,11 USD: hoa don 1.436.594, so 1.436.620,50.
            con_in = (tien - da_thu) if (nhan and d["tte"]) else con
            # HOA DON > (da thu + con lai): phan chenh la tien khach tra o KY TRUOC. PIAS da
            # tat dong do khoi bao cao cong no nen cot 'da thu' ky nay bang 0, de trong thi
            # khach doc thu khong hieu vi sao hoa don 363.875.386 ma chi doi 181.937.693. Vd
            # 26/45/08/AUTO/BE0016025: 2 hoa don cung ngay 30/06 moi cai 181.937.693, so chi
            # con 1 ky -> ghi 181.937.693 vao 'da thanh toan' cho 3 cot cong khop.
            dt_in = da_thu
            if nhan and tien > 0 and (tien - da_thu - con_in) >= 1:
                dt_in = tien - con_in
            c_tt = ws.cell(rr, 6, dt_in if abs(dt_in) >= 0.005 else None)
            c_tt.number_format = ws.cell(rr, 4).number_format
            ws.cell(rr, 7, con_in)
            # can giua theo chieu doc ca dong: style goc cua mau de mot so cot canh duoi
            # nen so don BH tut xuong day o trong khi STT va so hoa don nam giua
            for c in range(1, 8):
                a = ws.cell(rr, c).alignment
                ws.cell(rr, c).alignment = Alignment(
                    horizontal=a.horizontal, vertical="center",
                    wrap_text=a.wrap_text, indent=a.indent)
            tong_hd += tien
            tong_con += con_in
            if not nhan:
                log_rows.append(("Không tìm thấy hóa đơn", so_don, "", ten_sh, "thư gửi KH",
                                 ten, phai_thu, ""))
            elif abs(tien - phai_thu) >= 1:
                # doi >= 1.000 d so voi so thi khong con la chenh ty gia / lam tron nua,
                # phai noi ro thu da in khac so o cho nao
                ghi = []
                if abs(con_in - con) >= 1000:
                    ghi.append("thư in còn nợ theo hóa đơn")
                if abs(dt_in - da_thu) >= 1000:
                    ghi.append("thư ghi phần chênh vào Đã thanh toán")
                log_rows.append(("Tiền hóa đơn lệch tiền phải thu"
                                 + (f" ({so_ngay} lần xuất)" if so_ngay > 1 else "")
                                 + (", " + "; ".join(ghi) if ghi else ""),
                                 so_don, nhan, ten_sh, "thư gửi KH", ten, phai_thu, tien))
            elif abs(dt_in) >= 0.005 and abs(tien - dt_in - con_in) >= 1:
                # 3 so IN RA thu khong cong khop nhau du hoa don = so phai thu: PIAS quy doi
                # 'da thu' va 'con lai' theo ty gia khac nhau, vd 25/45/06/CHBB/P000179 (USD)
                # lech 1.417,68 d. Chi soat khi co in so da thanh toan — con 58 dong da thu = 0
                # lech vai chuc dong (HHXK: 127.626 vs 127.705,50) thi bo qua cho do nhieu.
                log_rows.append(("Đã thanh toán + còn lại lệch tiền hóa đơn", so_don, nhan,
                                 ten_sh, "thư gửi KH", ten, da_thu + con, tien))

        shift = n - 1
        for r0 in range(25, KH_ROW_END + 1):
            rr = r0 + shift
            _chep_dong(ws, tpl, rr, r0)
            for c in range(1, 11):
                v = mau[r0][c - 1]
                if isinstance(v, str) and v.startswith("="):
                    v = None
                ws.cell(rr, c, _thay_ngay(v, ngay, 2026, "30/06/2026", "quý 02/2026"))
        rt = 25 + shift
        ws.cell(rt, 4, f"=SUM(D{KH_ROW_DATA}:D{KH_ROW_DATA + n - 1})")
        ws.cell(rt, 6, f"=SUM(F{KH_ROW_DATA}:F{KH_ROW_DATA + n - 1})")
        ws.cell(rt, 7, f"=SUM(G{KH_ROW_DATA}:G{KH_ROW_DATA + n - 1})")
        ws.cell(rt + 1, 2, doc_tien(tong_con))
        for m in merges:
            rng = CellRange(m)
            if rng.min_row >= 25:
                ws.merge_cells(start_row=rng.min_row + shift, start_column=rng.min_col,
                               end_row=rng.max_row + shift, end_column=rng.max_col)
        _sua_chieu_cao(ws, 7, KH_ROW_END + shift)

    wb.remove(tpl)
    out = Path(out_dir) / f"Thư đối chiếu công nợ gửi KH {ngay:%d.%m.%Y}.xlsx"
    wb.save(out)
    log(f"  thu gui KH: {len(theo_kh) - bo} thu"
        + (f" (bo {bo} khach theo tham so trong config)" if bo else "") + f" -> {out.name}")
    return out, log_rows


# ------------------------------------------------------------- ban doi chieu ---
def duoi_le(co_tk):
    """Hau to ten file P.KD khi chay le, de khong nham voi ban day du cung ky."""
    return "" if len(co_tk) == 2 else f" ({co_tk[0]})"


def build_pkd(bang, phong, ngay, out_dir, log=print):
    tpl_path = file_kem(TPL_PKD, DIR_TPL)
    if not tpl_path.exists():
        raise SystemExit(f"Khong thay file mau:\n{tpl_path}")
    wb = load_workbook(tpl_path)

    gom = defaultdict(lambda: defaultdict(float))       # phong -> (tk, ky) -> tien
    co_tk = ["0051" if s == "0051" else "13111" for s in bang]
    log_rows = []
    thieu = set()
    for sheet, rows in bang.items():
        tk = "0051" if sheet == "0051" else "13111"
        for r in rows:
            pb = norm(r[C_PHONG])
            p = phong.get(pb)
            if not p or not p["sheet"]:
                thieu.add(clean_name(r[C_PHONG]) or "(trống)")
                continue
            ng = r[C_NGAY_CAP]
            if not isinstance(ng, (datetime, date)):
                ky = (0, 0)
            elif p.get("gom", "").startswith("n"):        # gom theo NAM
                ky = (ng.year, 0)
            else:                                        # gom theo THANG
                ky = (ng.year, ng.month)
            v = r[C_CON_LAI] if isinstance(r[C_CON_LAI], (int, float)) else 0
            gom[p["sheet"]][(tk, ky)] += v
    for t in sorted(thieu):
        log_rows.append(("Phòng ban chưa khai trong config.xlsx", "", "", "", "P.KD", t, "", ""))

    ten_sheet = {p["sheet"]: p for p in phong.values() if p["sheet"]}
    # Khai ten sheet trong config nhung file mau KHONG co sheet do => phong day bien mat
    # khoi ban xac nhan ma khong ai biet. Phai them sheet vao file mau trong Templates\
    # (nhan ban 1 sheet co san, doi ten + doi ten phong/nguoi ky) roi chay lai.
    co_mau = {s.strip() for s in wb.sheetnames}
    for sn in sorted(gom):
        if sn in co_mau:
            continue
        p = ten_sheet.get(sn)
        if not p:
            continue
        if _sheet_phong_moi(wb, sn, p, ten_sheet, log_rows) is None:
            log_rows.append(("Thiếu sheet trong file mẫu Templates, phòng này KHÔNG được xuất",
                             "", "", sn, "P.KD", p.get("ten") or p.get("pias") or "",
                             sum(gom[sn].values()), ""))
        else:
            log(f"    tu tao sheet moi cho phong: {sn}")
    giu = []
    for sn in list(wb.sheetnames):
        p = ten_sheet.get(sn.strip())
        if not p or sn.strip() not in {k.strip() for k in gom}:
            continue
        ct = _cau_truc_pkd(wb[sn])
        if not ct:
            log_rows.append(("Không đọc được bố cục sheet mẫu", "", "", sn, "P.KD", "", "", ""))
            continue
        giu.append(sn)
        d1, t1, d2, t2, r_tc, r_bc = ct
        ws = wb[sn]
        data = gom[sn.strip()]

        # CHUP style / chieu cao / noi dung TRUOC khi dong vao sheet — neu doc style truc tiep
        # tu sheet trong luc ghi thi cac dong nguon da bi ghi de, mat font va vien.
        NC = 11
        SN_ST = {r: [ws.cell(r, c)._style for c in range(1, NC + 1)]
                 for r in range(d1, ws.max_row + 1)}
        SN_H = {r: ws.row_dimensions[r].height for r in range(d1, ws.max_row + 1)}
        mau = {r: [ws.cell(r, c).value for c in range(1, NC + 1)]
               for r in range(d1, ws.max_row + 1)}
        merges = [str(m) for m in ws.merged_cells.ranges]

        def dat_style(r_dich, r_nguon, cot_a_tu=None):
            for c in range(1, NC + 1):
                nguon = cot_a_tu if (c == 1 and cot_a_tu) else r_nguon
                ws.cell(r_dich, c)._style = SN_ST[nguon][c - 1]
            if SN_H.get(r_nguon):
                ws.row_dimensions[r_dich].height = SN_H[r_nguon]

        def kem_ky(mau_d, ky):
            """Cot D cua sheet mau ghi kem ky CUA KY TRUOC — 'Hoa don theo bang ke dinh kem
            nam 2025' o TBHH/DTCN/PTKD, 'thang 03/2025' o PTSP lap lai cho ca 5 dong. Thay
            bang dung ky cua dong nay; mau nao khong kem ky (Bancass, KHDN) thi giu nguyen."""
            s = str(mau_d or "")
            if not s.strip():
                return mau_d
            if not ky[0]:
                return re.sub(r"\s*(năm \d{4}|tháng \d{1,2}/\d{4})\s*$", "", s) or s
            moi = f"tháng {ky[1]:02d}/{ky[0]}" if ky[1] else f"năm {ky[0]}"
            s2, n = re.subn(r"(năm \d{4}|tháng \d{1,2}/\d{4})\s*$", moi, s)
            return s2 if n else s

        def ky_chu(ky, mau_c):
            """Giu nguyen cach hanh van cua tung sheet mau, chi thay phan ky."""
            if not ky[0]:
                lb = "Các đơn BH không xác định ngày cấp đơn"
            elif ky[1]:
                lb = f"Các đơn BH cấp tháng {ky[1]:02d}/{ky[0]}"
            else:
                lb = f"Các đơn BH cấp năm {ky[0]}"
            duoi = str(mau_c or "").split("\n", 1)
            return lb + ("\n" + duoi[1] if len(duoi) > 1 else "")

        khoi = {tk: sorted((k, v) for (t, k), v in data.items() if t == tk)
                for tk in co_tk}

        for r0 in range(1, d1):
            v = ws.cell(r0, 1).value
            if isinstance(v, str) and v.startswith("Để phục vụ"):
                ws.cell(r0, 1, v.replace("tháng 06/2025", f"tháng {ngay:%m/%Y}")
                        .replace("30/06/2025", f"{ngay:%d/%m/%Y}"))
            elif isinstance(v, str) and "TK 13111-0051" in v and len(co_tk) == 1:
                ws.cell(r0, 1, v.replace("TK 13111-0051", f"TK {co_tk[0]}"))   # chay le
        _don_khoi(ws, d1)

        r = d1
        vi_tri = {}
        for tk, nguon in (("13111", d1), ("0051", d2)):
            if tk not in co_tk:                 # bao cao chay le: bo han khoi tai khoan kia
                continue
            dau = r
            # phong khong phat sinh o tai khoan nay (hay gap voi phong moi mo) — van phai
            # co dong 0 cho du 2 khoi, nhung dung ghi "Cac don BH khong xac dinh ngay cap don"
            trong = not khoi[tk]
            muc = khoi[tk] or [((0, 0), 0.0)]
            for i, (ky, v) in enumerate(muc):
                dat_style(r, nguon)
                ws.cell(r, 1, tk if i == 0 else None)
                ws.cell(r, 2, i + 1)
                ws.cell(r, 3, "Không phát sinh trong kỳ" if trong else ky_chu(ky, mau[nguon][2]))
                ws.cell(r, 4, None if trong else kem_ky(mau[nguon][3], ky))
                ws.cell(r, 5, v)
                ws.cell(r, 8, f"=E{r}")
                r += 1
            dat_style(r, t1)
            ws.cell(r, 1, "                              Tổng cộng")
            ws.cell(r, 5, f"=SUM(E{dau}:E{r - 1})")
            ws.cell(r, 8, f"=SUM(H{dau}:H{r - 1})")
            vi_tri[tk] = (dau, r)
            ws.merge_cells(start_row=dau, start_column=1, end_row=r - 1, end_column=1)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            r += 1

        r_tong = r
        dat_style(r_tong, r_tc)
        ws.cell(r_tong, 1, "                            TỔNG CỘNG")
        ws.cell(r_tong, 5, "=" + "+".join(f"E{t}" for _d, t in vi_tri.values()))
        ws.cell(r_tong, 8, "=" + "+".join(f"H{t}" for _d, t in vi_tri.values()))
        ws.merge_cells(start_row=r_tong, start_column=1, end_row=r_tong, end_column=4)

        tong = sum(data.values())
        shift = (r_tong + 1) - r_bc
        for r0 in range(r_bc, max(mau) + 1):
            rr = r0 + shift
            dat_style(rr, r0)
            for c in range(1, NC + 1):
                v = mau[r0][c - 1]
                if isinstance(v, str) and v.startswith("="):
                    v = None
                if isinstance(v, str):
                    v = v.replace("30/06/2025", f"{ngay:%d/%m/%Y}") \
                         .replace("năm 2025", f"năm {ngay.year}")
                ws.cell(rr, c, v)
        # Vung in cua file mau dung theo so dong cua ky mau. Ky nay nhieu dong hon la in ra
        # bi cat mat duoi — Bancass T07/2026 co noi dung toi dong 49 ma vung in van A1:H39,
        # cut ca khoi chu ky. Phai keo lai theo dong cuoi that.
        ws.print_area = f"A1:H{max(mau) + shift}"
        ws.cell(r_bc + shift, 3, doc_tien(tong, "ngàn", "đồng./."))
        for m in merges:
            rng = CellRange(m)
            if rng.min_row >= r_bc:
                ws.merge_cells(start_row=rng.min_row + shift, start_column=rng.min_col,
                               end_row=rng.max_row + shift, end_column=rng.max_col)
        _sua_chieu_cao(ws, 1, max(mau) + shift, 8)
        _ap_config_vao_sheet(ws, p, sn.strip(), log_rows)

    for sn in list(wb.sheetnames):
        if sn not in giu:
            wb.remove(wb[sn])
    out = Path(out_dir) / f"BẢN ĐỐI CHIẾU CÔNG NỢ GỬI P.KD T{ngay:%m.%Y}{duoi_le(co_tk)}.xlsx"
    wb.save(out)
    log(f"  ban doi chieu P.KD: {len(giu)} phong -> {out.name}")
    return out, log_rows, gom, giu


# --------------------------------------------------------------------- log ---
def ghi_log(rows, out_dir, ngay, log=print):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cần kiểm tra"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"CÁC ĐIỂM CẦN KIỂM TRA LẠI, số liệu đến {ngay:%d/%m/%Y}"
    ws["A1"].font = Font(name="Times New Roman", size=13, bold=True)
    hdr = ["Loại lệch", "Số hợp đồng / GCN", "Số hóa đơn", "Tên sheet", "Tên file",
           "Khách hàng / Phòng", "Số tiền phải thu", "Tiền trên hóa đơn", "Chênh lệch"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(3, j, h)
        c.font = Font(name="Times New Roman", size=11, bold=True)
        c.fill = PatternFill("solid", start_color="DDEBF7")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(j)].width = [34, 30, 34, 26, 16, 44, 18, 18, 18][j - 1]
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v if v != "" else None)
            c.font = Font(name="Times New Roman", size=10)
            if j in (7, 8):
                c.number_format = '#,##0.00;(#,##0.00);"-"'
        pt, hd = row[6], row[7]
        if isinstance(pt, (int, float)) and isinstance(hd, (int, float)):
            c = ws.cell(i, 9, hd - pt)
            c.font = Font(name="Times New Roman", size=10)
            c.number_format = '#,##0.00;(#,##0.00);"-"'
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:I{max(3, 3 + len(rows))}"
    out = Path(out_dir) / f"Log đối chiếu {ngay:%d.%m.%Y}.xlsx"
    wb.save(out)
    log(f"  log: {len(rows):,} dong can kiem tra -> {out.name}")
    return out


def build_all(report, out_dir, src_dir=None, log=print, progress=None, lam=("kh", "pkd"),
              inv_path=None, bk_path=None):
    step = progress or (lambda p, m: None)
    report = Path(report)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", report.name)
    ngay = date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else date.today()
    log(f"Doc bao cao: {report.name} (den ngay {ngay:%d/%m/%Y})")

    step(10, "Đọc báo cáo công nợ")
    # bao cao chay le (chi 0051 hoac chi 13111) thieu sheet kia: ban doi chieu P.KD chi co khoi
    # cua tai khoan co mat, thu gui KH can BC 13111
    bang = {k: v for k in ("BC 13111", "0051") if (v := doc_sheet(report, k)) is not None}
    if not bang:
        raise SystemExit(f"File bao cao khong co sheet 'BC 13111' hay '0051':\n{report}")
    for k, v in bang.items():
        log(f"  {k}: {len(v):,} dong")
    if "kh" in lam and "BC 13111" not in bang:
        raise SystemExit("Thu gui KH can TK 13111, ma bao cao nay chi co TK 0051.")

    inv = {}
    if "kh" in lam:
        step(25, "Đọc hóa đơn")
        # thu muc ky truoc, roi Input\, roi goc goi (goi cu de file hoa don canh .exe)
        thu_muc = (src_dir, report.parent, app_dir() / DIR_IN, app_dir())
        p = Path(inv_path) if inv_path else None
        if p and not p.exists():
            raise SystemExit(f"Khong thay file hoa don:\n{p}")
        p = p or tim_file(*thu_muc)
        if not p:
            raise SystemExit("Chua chon file 'List hóa đơn.xlsx', bam vao dong 'List hóa đơn' "
                             "de chon, hoac de file trong thu muc ky o Input\\.")
        log(f"  File hoa don: {p}")
        inv = doc_hoa_don(p, log, thu_muc, bk_path)

    kh_ma, kh_ten, phong, ts = doc_config()
    log(f"  config: {len(kh_ma):,} khach hang, {len(phong)} phong ban")
    log(f"  tham so lap thu: nguong={ts['nguong']:,.0f}, chi qua han={ts['chi_qua_han']}, "
        f"bo thu 0d={ts['bo_khong']}")

    f_kh = f_pkd = f_ann = None
    log1 = log2 = log3 = []
    if "kh" in lam:
        step(45, "Dựng thư gửi khách hàng")
        # thu gui KH chi lay TK 13111; 0051 la tai khoan noi bo (hop dong/tam ung),
        # chi dua vao ban doi chieu gui P.KD
        f_kh, log1 = build_kh({"BC 13111": bang["BC 13111"]}, inv, kh_ma, kh_ten, ts,
                              ngay, out_dir, log)
    if "pkd" in lam:
        step(75, "Dựng bản đối chiếu P.KD")
        # 0051: bo so don goc co tong con lai = 0 ('Tong cong no don BH bang 0') o CA ban xac
        # nhan lan bang ke. Chi bo o bang ke thi cap don goc cap thang 1 + SDBS cap thang 3
        # van nam tren ban xac nhan -> 2 file lech tung o (tai khoan x ky).
        bang_pkd = dict(bang)
        if "0051" in bang:
            rows51, bo_dong, bo_don = loc_0051(bang["0051"])
            if bo_dong:
                log(f"  0051: bo {bo_dong:,} dong cua {bo_don} so don goc co Check + - = 0")
            bang_pkd["0051"] = rows51
        f_pkd, log2, gom, giu = build_pkd(bang_pkd, phong, ngay, out_dir, log)
        # bang ke chi tiet dinh kem ban xac nhan — di lien voi ban doi chieu, 1 lan bam
        step(88, "Dựng bảng kê đính kèm")
        f_ann, log3 = build_annex(bang_pkd, phong, gom, giu, ngay, out_dir, log)
    step(95, "Ghi log")
    f_log = ghi_log(log1 + log2 + log3, out_dir, ngay, log)
    return f_kh, f_pkd, f_ann, f_log


if __name__ == "__main__":
    build_all(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else Path(sys.argv[1]).parent)
