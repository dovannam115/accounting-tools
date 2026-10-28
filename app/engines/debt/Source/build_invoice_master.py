"""Ghep 'BK_HDVAT_*.xlsx' (tien) voi 'List hoa don.xlsx' (so don BH) -> file trung gian.

Khoa ghep = Ky hieu hoa don + So hoa don (BK cot D+E  <->  List cot R 'Ky hieu HD' + B 'So ctu').
BK la nguon TIEN chuan: 1 hoa don xuat cho nhieu ky phi thi List chi gan duoc 1 ky,
vd BA-00005103 = 295.084.660 = 162.296.563 + 132.788.097 cua 26/45/06/BHKT/C000021.

Ra 2 sheet:
  'Hoa don'       1 dong = 1 hoa don  (tien BK, tien List, chenh lech, so don BH)
  'Don <-> Hoa don' 1 dong = 1 cap (so don BH x hoa don) de VLOOKUP theo so don

Cot 'Tien hoa don phan bo cho don' = tien BK chia lai cho tung don theo ty trong phi cua
List: hoa don 1 don thi an tron so BK (het lech voi don nhieu ky phi), hoa don gop nhieu
don thi chia dung ty le nen cong lai van bang BK. Ham phan_bo() la ban Python cua cong
thuc do — invoices.py dung ham nay chu KHONG doc gia tri cong thuc trong file.
"""
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from build_debt_report import DIR_IN, DIR_OUT

TNR = "Times New Roman"
F_DATA = Font(name=TNR, size=8)
F_HEAD = Font(name=TNR, size=9, bold=True)
F_TOTAL = Font(name=TNR, size=9, bold=True, color="C00000")
FILL_HEAD = PatternFill("solid", start_color="DDEBF7")
MONEY = '#,##0.00;(#,##0.00);"-"'
DATEF = "dd/mm/yyyy"

# List hoa don: ten header -> vi tri du phong (0-based)
COT_LIST = {"Số ctừ": 1, "Ngày ctừ": 2, "Số đơn Bh": 3, "Tk nợ": 4, "Khách": 13,
            "Phòng Kthác": 15, "Đại lý": 16, "Ký hiệu HĐ": 17, "Tiền phí": 26, "Tiền VAT": 28}


def so_hd(v):
    """Chuan hoa so hoa don ve chuoi 8 chu so.

    PIAS luc tra text luc tra so, va List con co ban 6-7 chu so:
    'BA 0000860' o List chinh la 'BA 00000860' o BK (cung ngay 10/02/2026).
    """
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, str):
        v = v.strip()
        if v.isdigit():
            v = int(v)
    return f"{v:08d}" if isinstance(v, int) else str(v)


def phan_bo(tien_don, tong_list, tong_bk, so_don):
    """Ban Python cua cong thuc cot I sheet 'Don <-> Hoa don'.

    tong_bk = None (hoa don khong co trong BK) -> giu nguyen tien theo List.
    Nguoc dau -> cung giu nguyen tien List: BA-00004589 la hoa don dieu chinh giam
    (BK = -107.019) nhung List gan vao do 2 don co phi DUONG (tong 873.716); chia ty
    trong se ra tien hoa don am cho don van dang phai thu.
    """
    if tong_bk is None:
        return tien_don
    if not tong_list:
        return tong_bk / (so_don or 1)
    if (tong_bk < 0) != (tong_list < 0):
        return tien_don
    return round(tong_bk * tien_don / tong_list, 2)


def doc_bk(path, log=print):
    """-> {(ky hieu, so hd): dict}, ky bao cao"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hd, ky, dau = {}, "", None
    for r in ws.iter_rows(values_only=True):
        if not r:
            continue
        if dau is None:
            txt = " ".join(str(v) for v in r if v)
            if "Từ ngày" in txt:
                ky = txt.strip()
            if "STT" in [str(v).strip() for v in r if v]:
                dau = True
            continue
        if len(r) < 12 or not r[4]:
            continue
        k = (str(r[3] or "").strip(), so_hd(r[4]))
        hd[k] = {"ngay": r[5], "ten": r[6], "mst": r[7], "truoc": r[8] or 0,
                 "thue": r[9] or 0, "tong": r[10] or 0, "tt": r[11],
                 "gc": r[12] if len(r) > 12 else None}
    wb.close()
    log(f"  {Path(path).name}: {len(hd):,} hóa đơn | {ky}")
    return hd, ky


def doc_list(path, log=print):
    """-> {(ky hieu, so hd): {so don BH: dict cong don}}"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(v or "").strip() for v in next(it)]
    ix = {t: (hdr.index(t) if t in hdr else d) for t, d in COT_LIST.items()}
    idx, n, bo = defaultdict(dict), 0, 0
    for r in it:
        if not r or len(r) <= ix["Tiền VAT"]:
            continue
        so = so_hd(r[ix["Số ctừ"]])
        don = str(r[ix["Số đơn Bh"]] or "").strip()
        if not so or not don:
            bo += 1
            continue
        k = (str(r[ix["Ký hiệu HĐ"]] or "").strip(), so)
        d = idx[k].get(don)
        if d is None:
            d = idx[k][don] = {"tien": 0.0, "ngay": r[ix["Ngày ctừ"]], "dong": 0,
                               "khach": r[ix["Khách"]], "phong": r[ix["Phòng Kthác"]],
                               "daily": r[ix["Đại lý"]], "tkno": r[ix["Tk nợ"]]}
        phi = r[ix["Tiền phí"]] if isinstance(r[ix["Tiền phí"]], (int, float)) else 0
        vat = r[ix["Tiền VAT"]] if isinstance(r[ix["Tiền VAT"]], (int, float)) else 0
        d["tien"] += phi + vat
        d["dong"] += 1
        if r[ix["Ngày ctừ"]] and (not d["ngay"] or r[ix["Ngày ctừ"]] > d["ngay"]):
            d["ngay"] = r[ix["Ngày ctừ"]]
        n += 1
    wb.close()
    log(f"  {Path(path).name}: {n:,} dòng, {len(idx):,} hóa đơn, "
        f"{sum(len(v) for v in idx.values()):,} cặp đơn×hóa đơn"
        + (f", bỏ {bo:,} dòng thiếu số ctừ/số đơn" if bo else ""))
    return dict(idx)


def _ke(ws, cols, r0):
    """Ke header + dinh dang cot; tra ve chi so dong du lieu dau tien."""
    for j, (ten, w, fmt) in enumerate(cols, 1):
        c = ws.cell(r0, j, ten)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(j)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{r0 + 1}"
    ws.row_dimensions[r0].height = 28
    return r0 + 1


def _do(ws, r, cols, vals):
    for j, ((_, _, fmt), v) in enumerate(zip(cols, vals), 1):
        c = ws.cell(r, j, v)
        c.font, c.number_format = F_DATA, fmt


def xay(bk, li, ky, out, log=print):
    wb = Workbook()

    # ------------------------------------------------------------ sheet 1 ---
    C1 = [("Ký hiệu", 8, "@"), ("Số hóa đơn", 12, "@"), ("Ngày lập HĐ", 12, DATEF),
          ("Tên khách hàng", 40, "@"), ("MST", 16, "@"),
          ("Tiền trước thuế", 16, MONEY), ("Tiền thuế", 14, MONEY),
          ("Tổng tiền HĐ\n(BK_HDVAT)", 17, MONEY), ("Tổng Phí+VAT\n(List hóa đơn)", 17, MONEY),
          ("Chênh lệch", 15, MONEY), ("Số đơn BH", 10, "0"), ("Số dòng List", 11, "0"),
          ("Tình trạng ghép", 18, "@"), ("Trạng thái HĐ", 14, "@"), ("Ghi chú (BK)", 26, "@")]
    ws = wb.active
    ws.title = "Hóa đơn"
    r = _ke(ws, C1, 2)
    r_dau = r

    for k in sorted(set(bk) | set(li), key=lambda k: ((bk.get(k) or {}).get("ngay")
                                                      or datetime(2100, 1, 1), k[0], k[1])):
        b, dons = bk.get(k), li.get(k, {})
        t_list = sum(d["tien"] for d in dons.values())
        if b and dons:
            tt = "Khớp" if round(b["tong"] - t_list, 2) == 0 else "Lệch tiền"
        elif b:
            tt = "Chỉ có ở BK"
        else:
            tt = "Chỉ có ở List"
        ngay = b["ngay"] if b else max((d["ngay"] for d in dons.values() if d["ngay"]),
                                       default=None)
        _do(ws, r, C1, [
            k[0], k[1], ngay,
            b["ten"] if b else (next(iter(dons.values()))["khach"] if dons else None),
            b["mst"] if b else None,
            b["truoc"] if b else None, b["thue"] if b else None, b["tong"] if b else None,
            t_list if dons else None,
            f"=IF(OR(H{r}=\"\",I{r}=\"\"),\"\",H{r}-I{r})",
            len(dons) or None, sum(d["dong"] for d in dons.values()) or None,
            tt, b["tt"] if b else None, b["gc"] if b else None])
        r += 1
    r_cuoi = r - 1
    for j, f in ((6, "F"), (7, "G"), (8, "H"), (9, "I"), (10, "J"), (11, "K"), (12, "L")):
        c = ws.cell(1, j, f"=SUBTOTAL(109,{f}{r_dau}:{f}{r_cuoi})")
        c.font, c.number_format = F_TOTAL, MONEY if j <= 10 else "#,##0"
    c = ws.cell(1, 1, "TỔNG")
    c.font = F_TOTAL
    ws.auto_filter.ref = f"A2:{L(len(C1))}{r_cuoi}"
    n_hd = r_cuoi - r_dau + 1

    # ------------------------------------------------------------ sheet 2 ---
    C2 = [("Số đơn BH", 24, "@"), ("Ký hiệu", 8, "@"), ("Số hóa đơn", 12, "@"),
          ("Ngày lập HĐ", 12, DATEF), ("Ngày ctừ (List)", 13, DATEF),
          ("Phí+VAT của đơn\ntrong hóa đơn", 18, MONEY),
          ("Σ Phí+VAT\ntrong hóa đơn", 18, MONEY),
          ("Tổng tiền HĐ\n(BK_HDVAT)", 17, MONEY),
          ("Tiền hóa đơn\nphân bổ cho đơn", 18, MONEY),
          ("Số đơn\ntrong HĐ", 10, "0"), ("Số dòng List", 11, "0"),
          ("Tên khách (List)", 34, "@"), ("Phòng khai thác", 14, "@"),
          ("Đại lý", 14, "@"), ("TK nợ", 9, "@"), ("Tình trạng ghép", 16, "@")]
    ws2 = wb.create_sheet("Đơn ↔ Hóa đơn")
    r = _ke(ws2, C2, 2)
    r_dau2 = r
    cap = [(don, k, d) for k, dons in li.items() for don, d in dons.items()]
    cap.sort(key=lambda x: (x[0], x[2]["ngay"] or datetime(2100, 1, 1), x[1][0], x[1][1]))
    for don, k, d in cap:
        b = bk.get(k)
        _do(ws2, r, C2, [
            don, k[0], k[1], b["ngay"] if b else None, d["ngay"], d["tien"],
            sum(x["tien"] for x in li[k].values()), b["tong"] if b else None,
            f'=IF(H{r}="",F{r},IF(G{r}=0,H{r}/J{r},'
            f'IF(SIGN(H{r})*SIGN(G{r})<0,F{r},ROUND(H{r}*F{r}/G{r},2))))',
            len(li[k]), d["dong"], d["khach"], d["phong"],
            d["daily"], str(d["tkno"] or ""), "Khớp" if b else "Không có ở BK"])
        r += 1
    r_cuoi2 = r - 1
    for j, f in ((6, "F"), (8, "H"), (9, "I"), (10, "J"), (11, "K")):
        c = ws2.cell(1, j, f"=SUBTOTAL(109,{f}{r_dau2}:{f}{r_cuoi2})")
        c.font, c.number_format = F_TOTAL, MONEY if j <= 9 else "#,##0"
    c = ws2.cell(1, 1, f"TỔNG, {ky}" if ky else "TỔNG")
    c.font = F_TOTAL
    ws2.auto_filter.ref = f"A2:{L(len(C2))}{r_cuoi2}"

    wb.save(out)
    log(f"  -> {Path(out).name}: {n_hd:,} hóa đơn / {len(cap):,} cặp đơn×hóa đơn")
    return n_hd, len(cap)


def nguon(goc):
    """Cac cho co the dat file export PIAS: tung thu muc ky trong Input\\, roi Input\\,
    cuoi cung la goc goi (goi cu de file ngay canh .exe)."""
    ra, inp = [], Path(goc) / DIR_IN
    if inp.is_dir():
        ra += [d for d in inp.iterdir() if d.is_dir()]
        ra.append(inp)
    ra.append(Path(goc))
    return ra


def tim(thu_muc, mau):
    """File moi nhat khop mau, quet het cac thu muc nguon."""
    ds = [thu_muc] if isinstance(thu_muc, (str, Path)) else list(thu_muc)
    fs = [p for d in ds for p in Path(d).glob(mau) if p.is_file()]
    return max(fs, key=lambda p: p.stat().st_mtime) if fs else None


def ky_thanh_nhan(ky):
    """'Tu ngay : 01/01/2026  Den ngay : 31/07/2026' -> '31.07.2026'"""
    d = re.search(r"Đến ngày\s*:\s*(\d{2})/(\d{2})/(\d{4})", ky or "")
    return f"{d.group(1)}.{d.group(2)}.{d.group(3)}" if d else datetime.now().strftime("%d.%m.%Y")


def dung(goc, log=print, f_bk=None, f_li=None, progress=None):
    """Ghep BK + List -> output\\Hoa don sach <ngay>.xlsx. Tra ve duong dan file."""
    step = progress or (lambda p, m: None)
    goc = Path(goc)
    tim_o = nguon(goc)
    f_bk = Path(f_bk) if f_bk else tim(tim_o, "BK_HDVAT*.xlsx")
    f_li = Path(f_li) if f_li else tim(tim_o, "List hóa đơn*.xlsx")
    if not f_bk or not f_li:
        raise SystemExit("Thieu file BK_HDVAT*.xlsx hoac List hoa don*.xlsx.\n"
                         f"De ca hai vao thu muc ky trong {goc / DIR_IN} roi bam lai.")
    step(15, "Đọc BK_HDVAT")
    bk, ky = doc_bk(f_bk, log)
    step(40, "Đọc List hóa đơn")
    li = doc_list(f_li, log)
    step(70, "Ghép và ghi file")
    out = goc / DIR_OUT / f"Hóa đơn sạch {ky_thanh_nhan(ky)}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    xay(bk, li, ky, out, log)
    return out


def main(argv):
    goc = Path(argv[1]) if len(argv) > 1 else Path(__file__).resolve().parent.parent
    dung(goc)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
