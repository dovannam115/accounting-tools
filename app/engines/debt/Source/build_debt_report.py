"""Dung file "BC cong no 1311" tu 5 file export PIAS (BCPBHTT_*.xlsx).

    py -3.13 build_debt_report.py <thu_muc_raw> [thu_muc_ghi_ket_qua]
"""
import re
import shutil
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

CONFIG_NAME = "config.xlsx"
# thu muc con cua goi ban giao — moi loai file mot cho, khong de tran o goc
DIR_IN, DIR_OUT, DIR_CFG, DIR_TPL = "Input", "Output", "Config", "Templates"
NHOM_MAC_DINH = "Nghiệp vụ khác"
# ma khach hang chung, nhin ten khong biet don cua ai: 3 sheet tong hop tach them dong con
# theo cot 'Ten moi gioi/ Ten dai ly' ngay duoi dong khach hang (yeu cau 14/09/2026)
KH_TACH_DAI_LY = ("Khách lẻ - Bán cho người tiêu dùng",)
DAI_LY_TRONG = "(Không ghi tên đại lý)"
# so don GOC (bo SDBS) co tong con lai = 0 thi khong con la cong no, bo khoi bao cao qua han
TT_KHONG_NO = "Tổng công nợ đơn BH bằng 0"
# cau truc co dinh cua bao cao, khong de nguoi dung sua
ACCOUNTS = {"131112": "12", "131113": "13", "131114": "14", "131115": "15", "0051": "0051"}
MERGE_SHEET, MERGE_FROM = "BC 13111", ["12", "13", "14", "15"]
# chay le (30/09/2026, ke toan muon check rieng 0051 theo phong KD ma khong dung 13111):
# build(phan=("0051",)) chi can file 0051, file tai khoan khac co trong thu muc thi bo qua
PHAN_DU = ("13111", "0051")
TEN_BC = {PHAN_DU: "BC công nợ 1311", ("13111",): "BC công nợ 13111", ("0051",): "BC công nợ 0051"}


def phan_cua(acc):
    return "0051" if acc == "0051" else "13111"


def app_dir():
    """Thu muc goc cua goi: canh file .exe khi da dong goi, canh tool\\ khi chay .py."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


def file_kem(ten, con):
    """File di kem goi (config, file mau): trong thu muc con truoc, roi den goc goi.
    Van doc duoc goi cu chua tach thu muc."""
    for p in (app_dir() / con / ten, app_dir() / ten):
        if p.exists():
            return p
    return app_dir() / con / ten          # chua co that -> bao loi theo duong dan chuan


def config_path():
    """config.xlsx trong Config\\ cho de sua; thieu thi lay ban mac dinh kem trong goi."""
    cfg = file_kem(CONFIG_NAME, DIR_CFG)
    if not cfg.exists():
        bundled = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / CONFIG_NAME
        if not bundled.exists():
            raise SystemExit(f"Khong tim thay {CONFIG_NAME} o:\n{cfg}")
        cfg.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(bundled, cfg)
    return cfg


def load_config():
    """Doc config.xlsx -> (nhom nghiep vu, thu tu khoi, luat gop khach hang)."""
    wb = load_workbook(config_path(), data_only=True)
    nhom, order = {}, []
    for r in wb["Nhóm nghiệp vụ"].iter_rows(min_row=3, values_only=True):
        ten, grp = (r + (None, None))[:2]
        if not ten or not grp:
            continue
        grp = str(grp).strip()
        nhom[norm(ten)] = grp
        if grp not in order and grp != NHOM_MAC_DINH:
            order.append(grp)
    gop = []
    for r in wb["Gộp khách hàng"].iter_rows(min_row=3, values_only=True):
        kieu, mau, dich, ghi = (tuple(r) + (None,) * 4)[:4]
        if not mau or not dich:
            continue
        gop.append((str(kieu or "Chứa chữ").strip(), str(mau).strip(),
                    clean_name(dich), str(ghi or "").strip()))
    wb.close()
    if not nhom:
        raise SystemExit(f"Sheet 'Nhóm nghiệp vụ' trong {CONFIG_NAME} dang trong.")
    return nhom, order + [NHOM_MAC_DINH], gop

# ------------------------------------------------------------------ layout ---
# Vi tri cot trong file raw PIAS (0-based, dong 13 la ten field DB)
RAW = dict(stt=0, ma_kh=1, ten_kh=2, ma_cbkt=3, so_donbh=4, ten_cbcnv=5, ma_pban=6,
           ten_pban=7, ten_donbh=9, ma_kthac=10, ten_daily=11, ngay_dau=12, ngay_cuoi=13,
           ngay_donbh=14, ngay_ttoan=15, ma_tte=16, tygia=17, nte_ps=18, nte_tt=19, nte=20,
           tien_ps=21, tien_tt=22, tien=23, qua_han=24, no09=25, no18=26, no30=27, no90=28,
           no180=29, no200=30, dien_giai=33, ma_hdong=34, ttoan_old=35, ttoan_pias=36)

# 24 cot dau dung chung cho moi sheet du lieu
BASE24 = [
    ("STT", "stt", "int"), ("Mã khách hàng", "ma_kh", "text"),
    ("Tên khách hàng", "ten_kh", "text"), ("Mã CBCNV", "ma_cbkt", "text"),
    ("Số đơn bảo hiểm", "so_donbh", "text"), ("Tên CB", "ten_cbcnv", "text"),
    ("Mã Phòng ban", "ma_pban", "text"), ("Phòng ban", "ten_pban", "text"),
    ("Tên đơn BH", "ten_donbh", "text"), ("Mã khai thác", "ma_kthac", "int"),
    ("Tên môi giới/ Tên đại lý", "ten_daily", "text"),
    ("Thời hạn bảo hiểm Từ ngày", "ngay_dau", "date"), ("Đến ngày", "ngay_cuoi", "date"),
    ("Ngày cấp đơn", "ngay_donbh", "date"), ("Ngày thu phí", "ngay_ttoan", "date"),
    ("Loại tiền", "ma_tte", "text"), ("Tỷ giá", "tygia", "rate"),
    ("Nguyên tệ phải thu", "nte_ps", "money"), ("Nguyên tệ đã thu", "nte_tt", "money"),
    ("Nguyên tệ còn lại phải thu", "nte", "money"),
    ("Số tiền phải thu", "tien_ps", "money"), ("Số tiền đã thu", "tien_tt", "money"),
    ("Số tiền còn lại phải thu", "tien", "money"), ("Tổng số quá hạn", "qua_han", "money"),
]
BUCKETS = ["Quá hạn >=0 & <90 ngày", "Quá hạn >=90 & <180 ngày", "Quá hạn >=180 & <365 ngày",
           "Quá hạn >=1 & <2 năm", "Quá hạn >=2 & <3 năm", "Quá hạn >3 năm"]
BUCKET_BOUNDS = [(0, 90), (90, 180), (180, 365), (365, 730), (730, 1095), (1095, None)]

# sheet 13111x: 24 cot goc + 6 bucket cua PIAS + 4 cot phu
LAYOUT_13111 = BASE24 + [(b, k, "money") for b, k in
                         zip(BUCKETS, ["no09", "no18", "no30", "no90", "no180", "no200"])] + [
    ("Ghi chú", "dien_giai", "text"), ("Mã hợp đồng", "ma_hdong", "text"),
    ("Hạn t.toán khách hàng", "ttoan_old", "date"),
    ("Hạn t.toán nhập tay PIAS", "ttoan_pias", "date"),
]
# sheet 0051: 24 cot goc + so ngay qua han + 6 bucket TINH LAI + cot kiem tra
LAYOUT_0051 = BASE24 + [("Số ngày quá hạn", None, "int")] + \
    [(b, None, "money") for b in BUCKETS] + [
    ("Ghi chú", "dien_giai", "text"), ("Mã hợp đồng", "ma_hdong", "text"),
    ("Số đơn BH gốc (bỏ SDBS)", None, "text"), ("Check + -", None, "money"),
    ("TÌNH TRẠNG TT", None, "text"), ("KD kiểm tra", None, "text"),
    ("Cách xử lý\n(Hủy đơn, Hủy trắng, Xuất HĐ)", None, "text"),
]
EXTRA_13111 = [("Nhóm nghiệp vụ", "calc"), ("KH gộp", "calc"), ("KH chuẩn hoá", "calc"),
               ("Đại lý chuẩn hoá", "calc"),
               ("Nguyên Nhân", "input"), ("KH thanh toán/đối trừ", "input")]
EXTRA_0051 = [("KH chuẩn hoá", "calc"), ("Đại lý chuẩn hoá", "calc")]

ROW_TOTAL, ROW_HEAD, ROW_DATA = 10, 11, 12

NUMFMT = {"int": "0", "text": "@", "date": "dd/mm/yyyy", "rate": "#,##0",
          "money": '#,##0.00;(#,##0.00);"-"'}
WIDTH = {"int": 6, "text": 22, "date": 12, "rate": 9, "money": 16}

TNR = "Times New Roman"
F_DATA = Font(name=TNR, size=8)
F_HEAD = Font(name=TNR, size=9, bold=True)
F_TOTAL = Font(name=TNR, size=9, bold=True, color="C00000")
F_TITLE = Font(name=TNR, size=12, bold=True)
F_NOTE = Font(name=TNR, size=9, italic=True)
FILL_HEAD = PatternFill("solid", start_color="DDEBF7")
FILL_CALC = PatternFill("solid", start_color="EDEDED")
FILL_INPUT = PatternFill("solid", start_color="FFF2CC")
THIN = Side(style="thin", color="B0B0B0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ------------------------------------------------------------------- utils ---
def vn_key(s):
    """Khoa sap xep gan voi thu tu A->Z cua Excel tieng Viet."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().upper()


def clean_name(s):
    """Giong het TRIM(CLEAN(x)) cua Excel: bo ky tu dieu khien (co CHAR(10) ma PIAS
    hay de lan trong ten khach hang), gop khoang trang thua, cat 2 dau."""
    s = re.sub(r"[\x00-\x1F\x7F]", "", str(s or ""))
    return re.sub(r" +", " ", s).strip()


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def don_goc(so_don):
    """So don goc, bo duoi sua doi bo sung /E<so>: 26/45/09/KHAC/P000650/E04 -> .../P000650.
    Duoi toan so (/13, /15) khong phai SDBS nen giu nguyen. Phai khop f_don_goc()."""
    s = str(so_don or "").strip()
    return re.sub(r"/E\d{1,4}$", "", s) if s.count("/") >= 5 else s


def f_don_goc(r):
    """Cong thuc Excel y het don_goc() cho dong r, doc cot E. Khong dung LET/TEXTBEFORE de
    chay duoc ca Excel 2010-2019."""
    e = f"$E{r}"
    k = f'FIND("/E",{e},MAX(1,LEN({e})-5))'
    return (f'=IFERROR(IF(AND(LEN({e})-LEN(SUBSTITUTE({e},"/",""))>=5,'
            f'ISERROR(FIND("/",{e},{k}+1)),ISNUMBER(-MID({e},{k}+2,9))),'
            f'LEFT({e},{k}-1),{e}&""),{e}&"")')


def tieu_chi(ten):
    """Tieu chi SUMIFS khop dung y ten (thoat * ? ~ va dau nhay). Ten trong -> "" khop o
    trong lan chuoi rong do TRIM(CLEAN()) tra ve."""
    if not ten:
        return '""'
    ten = ten.replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')
    return f'"={ten}"'


def dai_ly_theo(cap):
    """cap: (khoa, ten dai ly) -> {khoa: [ten dai ly da sap]}. SUMIFS khong phan biet hoa
    thuong nen gom theo chu hoa, giu cach viet gap dau tien; dai ly trong xep cuoi."""
    gom = {}
    for khoa, dl in cap:
        dl = clean_name(dl)
        gom.setdefault(khoa, {}).setdefault(dl.upper(), dl)
    return {k: sorted(v.values(), key=lambda s: (s == "", vn_key(s))) for k, v in gom.items()}


def can_tach(kh):
    return norm(kh) in {norm(k) for k in KH_TACH_DAI_LY}


def tru_dong_con(col, rows):
    """'-SUM(B52:B57)...' cho tung day dong con lien nhau: dong Tong cong chi cong dong cha."""
    out, dau = "", None
    for i, r in enumerate(rows):
        dau = r if dau is None else dau
        if i + 1 == len(rows) or rows[i + 1] != r + 1:
            out += f"-SUM({col}{dau}:{col}{r})"
            dau = None
    return out


def ghi_dong_con(ws, r, ten, font):
    c = ws.cell(r, 1, ten or DAI_LY_TRONG)
    c.font, c.alignment = font, Alignment(indent=2)


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for f in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def export_ts(path):
    """Phan duoi ten file BCPBHTT_<DDMMYYYY> <HHMMSS> chi la thoi diem xuat bao cao.
    Ma tai khoan KHONG nam trong ten file, phai doc o dong 6 ben trong file."""
    m = re.search(r"BCPBHTT[_\- ]*(\d{2})(\d{2})(\d{4})[ _\-]*(\d{2})(\d{2})(\d{2})", path.name)
    if m:
        d, mo, y, hh, mi, ss = (int(x) for x in m.groups())
        try:
            return datetime(y, mo, d, hh, mi, ss)
        except ValueError:
            pass
    return datetime.fromtimestamp(path.stat().st_mtime)


def read_raw(path):
    """Doc 1 file export PIAS -> (ma tai khoan, ngay bao cao, tieu de, list dong)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    head, rows = {}, []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 13:
            head[i] = row
            continue
        if row is None or len(row) <= RAW["ma_kh"] or row[RAW["ma_kh"]] in (None, ""):
            continue
        rows.append(row)
    wb.close()

    def cell(r, c=3):
        v = head.get(r) or ()
        return v[c] if len(v) > c else None

    acc = None
    for c in range(0, 8):
        m = re.search(r":\s*([0-9]{4,6})", str(cell(6, c) or ""))
        if m:
            acc = m.group(1)
            break
    rdate = None
    for c in range(0, 8):
        rdate = as_date(cell(5, c))
        if rdate:
            break
    title = next((v for v in (head.get(4) or ()) if v), "BÁO CÁO PHẢI THU PHÍ BẢO HIỂM")
    return acc, rdate, str(title).strip(), rows


def append_new_products(names, log=print):
    """Nghiep vu moi phat sinh -> them vao cuoi sheet 'Nhom nghiep vu' de ky sau nhin thay."""
    try:
        cfg = config_path()
        wb = load_workbook(cfg)
        ws = wb["Nhóm nghiệp vụ"]
        r = ws.max_row + 1
        for n in names:
            src = ws.cell(r - 1, 1)
            for j, v in enumerate((n, NHOM_MAC_DINH), 1):
                c = ws.cell(r, j, v)
                c.font, c.border, c.fill = src.font.copy(), src.border.copy(), src.fill.copy()
            r += 1
        wb.save(cfg)
        wb.close()
        log(f"  da them {len(names)} nghiep vu moi vao {CONFIG_NAME} (mac dinh '{NHOM_MAC_DINH}')")
    except PermissionError:
        log(f"  ! {CONFIG_NAME} dang mo trong Excel nen khong ghi them duoc, dong file roi chay lai")
    except Exception as e:                                   # noqa: BLE001
        log(f"  ! khong ghi duoc vao {CONFIG_NAME}: {e}")


NHOM, GROUPS, GOP = {}, [NHOM_MAC_DINH], []      # nap tu config.xlsx khi chay build()


def resolve_group(ten_donbh):
    return NHOM.get(norm(ten_donbh), NHOM_MAC_DINH)


def resolve_customer(ten_kh):
    """Xet luat gop tu tren xuong, dong nao khop truoc thi lay dong do."""
    ten = clean_name(ten_kh)
    for kieu, mau, dich, ghi in GOP:
        k = kieu.lower()
        if k.startswith("đúng") or k.startswith("dung"):
            hit = norm(ten) == norm(mau)
        elif k.startswith("regex"):
            hit = re.search(mau, ten, re.IGNORECASE) is not None
        else:                                     # "Chua chu" — bo dau, khong phan biet hoa thuong
            hit = vn_key(mau) in vn_key(ten)
        if hit:
            return dich, ghi
    return ten, ""


# ------------------------------------------------------------- sheet writer ---
def write_header_block(ws, title, acc_label, rdate, ncols):
    ws.sheet_view.showGridLines = False
    ws["A4"], ws["A4"].font = title, F_TITLE
    for r, (lbl, val) in enumerate(
            [("Đến ngày", rdate), ("Danh mục tài khoản", acc_label),
             ("Danh mục khách hàng", "Tất cả"), ("Danh mục phòng ban", "Tất cả"),
             ("Danh mục sản phẩm", "Tất cả")], start=5):
        ws.cell(r, 1, lbl).font = Font(name=TNR, size=9)
        c = ws.cell(r, 4, val)
        c.font = Font(name=TNR, size=9, bold=True)
        if isinstance(val, date):
            c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = f"A{ROW_DATA}"
    ws.auto_filter.ref = f"A{ROW_HEAD}:{L(ncols)}{ROW_HEAD}"


def write_data_sheet(wb, name, layout, rows, title, acc_label, rdate, extra=()):
    ws = wb.create_sheet(name)
    ncols = len(layout) + len(extra)
    write_header_block(ws, title, acc_label, rdate, ncols)

    for i, (hdr, _key, kind) in enumerate(layout, 1):
        c = ws.cell(ROW_HEAD, i, hdr)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(i)].width = WIDTH[kind]
    for j, (hdr, role) in enumerate(extra, len(layout) + 1):
        c = ws.cell(ROW_HEAD, j, hdr)
        c.font, c.border = F_HEAD, BOX
        c.fill = FILL_CALC if role == "calc" else FILL_INPUT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(j)].width = 24
    ws.row_dimensions[ROW_HEAD].height = 30

    for r, src in enumerate(rows, ROW_DATA):
        for i, (_h, key, kind) in enumerate(layout, 1):
            c = ws.cell(r, i)
            c.font, c.number_format = F_DATA, NUMFMT[kind]
            if key is None:
                continue
            v = src[RAW[key]] if len(src) > RAW[key] else None
            if kind == "date":
                v = as_date(v)
            elif kind == "text" and v is not None:
                v = str(v)
            c.value = v
    return ws


def add_total_row(ws, cols, last):
    ws.cell(ROW_TOTAL, 1, "TỔNG").font = F_TOTAL
    for i in cols:
        c = ws.cell(ROW_TOTAL, i, f"=SUM({L(i)}{ROW_DATA}:{L(i)}{last})")
        c.font, c.number_format = F_TOTAL, NUMFMT["money"]
        c.border = BOX


def ghi_qua_han(wb, ten, tieu_de, ghi_chu, nguon, last, c_kh, c_dl, cols, khs, con, them=""):
    """Sheet tong hop theo khach hang bang SUMIFS tro ve sheet nguon. Khach can tach co dong
    con theo dai ly ngay duoi; cuoi sheet co dong KIEM TRA voi dong TONG cua sheet nguon.
    cols: [(tieu de, cot nguon, co kiem tra khong)]; them: dieu kien SUMIFS bo sung."""
    ws = wb.create_sheet(ten)
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A1"].font = tieu_de, F_TITLE
    if ghi_chu:
        ws["A2"], ws["A2"].font = ghi_chu, F_NOTE
    for j, hdr in enumerate(["Khách hàng"] + [h for h, _s, _k in cols], 1):
        c = ws.cell(3, j, hdr)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[L(j)].width = 46 if j == 1 else 17
    ws.row_dimensions[3].height = 30

    def rng(c):
        return f"'{nguon}'!${c}${ROW_DATA}:${c}${last}"

    f_con = Font(name=TNR, size=8, italic=True)
    r, dong_con = 4, []
    for kh in khs:
        cha = r
        ws.cell(r, 1, kh).font = F_DATA
        for j, (_h, sc, _k) in enumerate(cols, 2):
            c = ws.cell(r, j, f"=SUMIFS({rng(sc)},{rng(c_kh)},$A{r}{them})")
            c.font, c.number_format = F_DATA, NUMFMT["money"]
        r += 1
        for dl in con.get(kh, []):
            ghi_dong_con(ws, r, dl, f_con)
            for j, (_h, sc, _k) in enumerate(cols, 2):
                c = ws.cell(r, j, f"=SUMIFS({rng(sc)},{rng(c_kh)},$A${cha},"
                                  f"{rng(c_dl)},{tieu_chi(dl)}{them})")
                c.font, c.number_format = f_con, NUMFMT["money"]
            dong_con.append(r)
            r += 1
    tong = r
    ws.cell(tong, 1, "Tổng cộng").font = F_TOTAL
    for j in range(2, len(cols) + 2):
        c = ws.cell(tong, j, f"=SUM({L(j)}4:{L(j)}{tong - 1}){tru_dong_con(L(j), dong_con)}")
        c.font, c.number_format, c.border = F_TOTAL, NUMFMT["money"], BOX
    ws.cell(tong + 2, 1, f"KIỂM TRA: Tổng cộng − dòng TỔNG sheet {nguon} (phải = 0)").font = F_TOTAL
    for j, (_h, sc, kiem) in enumerate(cols, 2):
        if kiem:
            c = ws.cell(tong + 2, j, f"={L(j)}{tong}-'{nguon}'!${sc}${ROW_TOTAL}")
            c.font, c.number_format, c.border = F_TOTAL, NUMFMT["money"], BOX
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{L(len(cols) + 1)}{max(tong - 1, 4)}"
    return ws


# -------------------------------------------------------------------- build ---
def build(src_dir, out_dir=None, log=print, progress=None, phan=PHAN_DU):
    """Doc raw trong src_dir -> ghi file bao cao vao out_dir. Tra ve duong dan file.
    phan: ("13111", "0051") bao cao day du; ("0051",) hay ("13111",) de chay le."""
    global NHOM, GROUPS, GOP
    phan = tuple(p for p in PHAN_DU if p in phan)
    if not phan:
        raise SystemExit("Chua chon tai khoan nao (13111, 0051) de dung bao cao.")
    co13, co51 = "13111" in phan, "0051" in phan
    src_dir = Path(src_dir).resolve()
    out_dir = Path(out_dir).resolve() if out_dir else src_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    step = (lambda p, m: progress(p, m)) if progress else (lambda p, m: None)

    NHOM, GROUPS, GOP = load_config()
    log(f"Doc {CONFIG_NAME}: {len(NHOM)} nghiep vu, {len(GOP)} luat gop khach hang")
    log(f"  cac khoi se dung: {' | '.join(GROUPS)}")

    files = sorted(src_dir.glob("BCPBHTT_*.xlsx"))
    if not files:
        raise SystemExit(f"Khong thay file BCPBHTT_*.xlsx nao trong:\n{src_dir}")

    step(5, "Doc file export PIAS")
    log(f"Doc raw tu: {src_dir}")
    if phan != PHAN_DU:
        log(f"  CHAY LE: chi TK {' + '.join(phan)}")
    by_acc, rdates = {}, []
    for f in sorted(files, key=export_ts):
        acc, rdate, title, rows = read_raw(f)
        ts = export_ts(f)
        if not acc:
            log(f"  ! bo qua {f.name}: khong doc duoc ma tai khoan o dong 6")
            continue
        if acc in ACCOUNTS and phan_cua(acc) not in phan:
            log(f"  - bo qua {f.name}: TK {acc} khong nam trong lan chay le nay")
            continue
        if acc in by_acc and by_acc[acc]["ts"] >= ts:
            log(f"  ! {f.name}: da co ban xuat moi hon cho TK {acc}, bo qua")
            continue
        if acc in by_acc:
            log(f"  ! TK {acc}: dung ban xuat moi hon ({f.name}), bo {by_acc[acc]['name']}")
        by_acc[acc] = dict(rows=rows, title=title, rdate=rdate, ts=ts, name=f.name)
        rdates.append(rdate)
        log(f"  {f.name:<34} TK {acc:<6} {len(rows):>6} dong  den ngay {rdate}"
            f"  (xuat {ts:%d/%m/%Y %H:%M})")

    want = {a: s for a, s in ACCOUNTS.items() if phan_cua(a) in phan}
    missing = [a for a in want if a not in by_acc]
    if missing:
        raise SystemExit(f"Thieu file export cho tai khoan: {', '.join(missing)}")
    if len({d for d in rdates if d}) > 1:
        log(f"  ! CANH BAO: cac file khac ngay bao cao: {sorted({str(d) for d in rdates if d})}")
    rdate = max(d for d in rdates if d)

    wb = Workbook()
    wb.remove(wb.active)

    if co13:
        # --- sheet tung tai khoan con 13111x -------------------------------------
        step(20, "Dung sheet 12 / 13 / 14 / 15")
        union_rows = []
        for acc, sheet in want.items():
            if acc == "0051":
                continue
            d = by_acc[acc]
            ws = write_data_sheet(wb, sheet, LAYOUT_13111, d["rows"], d["title"],
                                  f"{acc} – Phải thu phí bảo hiểm", rdate)
            add_total_row(ws, range(18, 31), ROW_DATA + len(d["rows"]) - 1)
            if sheet in MERGE_FROM:
                union_rows.extend(d["rows"])

        # --- sheet hop nhat BC 13111 --------------------------------------------
        step(45, "Hop nhat sheet BC 13111")
        bc_name = MERGE_SHEET
        d0 = by_acc[next(iter(want))]
        bc = write_data_sheet(wb, bc_name, LAYOUT_13111, union_rows, d0["title"],
                              "13111 – Phải thu phí bảo hiểm (hợp nhất)", rdate, extra=EXTRA_13111)
        bc_last = ROW_DATA + len(union_rows) - 1
        add_total_row(bc, range(18, 31), bc_last)
        n13111 = len(LAYOUT_13111)
        c_grp, c_kh, c_std, c_dl = n13111 + 1, n13111 + 2, n13111 + 3, n13111 + 4
        L_KH, L_STD, L_GRP, L_DL = L(c_kh), L(c_std), L(c_grp), L(c_dl)
        for r in range(ROW_DATA, bc_last + 1):
            cells = (
                bc.cell(r, c_grp, f'=IFERROR(VLOOKUP($I{r},Config!$A$4:$B$999,2,FALSE),"Nghiệp vụ khác")'),
                bc.cell(r, c_kh, f"=IFERROR(VLOOKUP(${L_STD}{r},Config!$D$4:$E$999,2,FALSE),${L_STD}{r})"),
                bc.cell(r, c_std, f"=TRIM(CLEAN($C{r}))"),
                bc.cell(r, c_dl, f"=TRIM(CLEAN($K{r}))"),
            )
            for c in cells:
                c.font, c.number_format, c.fill = F_DATA, "@", FILL_CALC

    if co51:
        # --- sheet 0051 ----------------------------------------------------------
        step(60, "Dung sheet 0051 + tinh lai tuoi no")
        d51 = by_acc["0051"]
        ws51 = write_data_sheet(wb, want["0051"], LAYOUT_0051, d51["rows"], d51["title"],
                                "0051 – Hợp đồng/tạm ứng", rdate, extra=EXTRA_0051)
        last51 = ROW_DATA + len(d51["rows"]) - 1
        cY = len(BASE24) + 1                      # So ngay qua han
        cZ = cY + 1                               # bucket dau tien
        cGoc, cChk, cTT = cY + 9, cY + 10, cY + 11          # AH / AI / AJ
        c51_kh, c51_dl = len(LAYOUT_0051) + 1, len(LAYOUT_0051) + 2
        for r in range(ROW_DATA, last51 + 1):
            ws51.cell(r, cY, f'=IF($O{r}="","",$D$5-$O{r})').number_format = "0"
            for b, (lo, hi) in enumerate(BUCKET_BOUNDS):
                cond = (f"AND(ISNUMBER($Y{r}),$Y{r}>={lo},$Y{r}<{hi})" if hi
                        else f"AND(ISNUMBER($Y{r}),$Y{r}>={lo})")
                ws51.cell(r, cZ + b, f"=IF({cond},$X{r},0)")
            # Check + - cong theo so don GOC: don goc va cac SDBS /E01, /E02... am duong trieu
            # tieu nhau (vd P000650/E04 + P000650/E05) thi ra 0 -> khong con la cong no
            ws51.cell(r, cGoc, f_don_goc(r))
            ws51.cell(r, cChk,
                      f"=SUMIF($AH${ROW_DATA}:$AH${last51},$AH{r},$W${ROW_DATA}:$W${last51})")
            ws51.cell(r, cTT,
                      f'=IF(ABS($AI{r})<0.005,"{TT_KHONG_NO}",IF(NOT(ISNUMBER($Y{r})),"",'
                      f'IF($Y{r}<0,"Chưa đến hạn thanh toán","Quá hạn thanh toán")))')
            for c in range(cY, cTT + 1):
                ws51.cell(r, c).font = F_DATA
            for c, f in ((c51_kh, f"=TRIM(CLEAN($C{r}))"), (c51_dl, f"=TRIM(CLEAN($K{r}))")):
                x = ws51.cell(r, c, f)
                x.font, x.number_format, x.fill = F_DATA, "@", FILL_CALC
        add_total_row(ws51, [23, 24] + list(range(cZ, cZ + 6)) + [cChk], last51)
        for c in (cTT + 1, cTT + 2):
            ws51.cell(ROW_HEAD, c).fill = FILL_INPUT
        dv = DataValidation(type="list", formula1='"Hủy đơn,Hủy trắng,Xuất HĐ"', allow_blank=True)
        ws51.add_data_validation(dv)
        dv.add(f"{L(cTT + 2)}{ROW_DATA}:{L(cTT + 2)}{last51}")

        # tinh san bang Python y het cong thuc AH/AI/AJ: dong nao vao 'Qua han 0051'
        tong51 = defaultdict(float)
        for r in d51["rows"]:
            v = r[RAW["tien"]]
            tong51[don_goc(r[RAW["so_donbh"]])] += v if isinstance(v, (int, float)) else 0
        rows51_no = [r for r in d51["rows"] if abs(tong51[don_goc(r[RAW["so_donbh"]])]) >= 0.005]
        log(f"  0051: {len(d51['rows']) - len(rows51_no):,} dong cua "
            f"{sum(1 for v in tong51.values() if abs(v) < 0.005):,} so don goc co tong con lai = 0"
            f" -> '{TT_KHONG_NO}', khong dua vao 'Quá hạn 0051'")

    if co13:
        # --- Config --------------------------------------------------------------
        step(72, "Ghi bang tham chieu Config")
        prods = sorted({str(r[RAW["ten_donbh"]] or "") for r in union_rows}, key=vn_key)
        moi = [p for p in prods if p and norm(p) not in NHOM]
        if moi:
            log(f"\n[NGHIEP VU MOI] chua khai trong {CONFIG_NAME}, tam xep vao '{NHOM_MAC_DINH}':")
            for p in moi:
                log(f"    {p}")
            append_new_products(moi, log)
        khs = sorted({clean_name(r[RAW["ten_kh"]]) for r in union_rows}, key=vn_key)
        gop = [(k,) + resolve_customer(k) for k in khs]
        gop = [g for g in gop if norm(g[0]) != norm(g[1])]
        cfg_ws = wb.create_sheet("Config")
        cfg_ws.sheet_view.showGridLines = False
        cfg_ws["A1"] = "Bảng tham chiếu"
        cfg_ws["A1"].font = F_TITLE
        for col, hdr in [("A", "Tên đơn BH"), ("B", "Nhóm nghiệp vụ"),
                         ("D", "Tên khách hàng gốc"), ("E", "Gộp thành"), ("F", "Ghi chú")]:
            c = cfg_ws[f"{col}3"]
            c.value, c.font, c.fill, c.border = hdr, F_HEAD, FILL_HEAD, BOX
            cfg_ws.column_dimensions[col].width = 44 if col in ("A", "D", "E") else 22
        for i, p in enumerate(prods, 4):
            cfg_ws.cell(i, 1, p).font = F_DATA
            cfg_ws.cell(i, 2, resolve_group(p)).font = F_DATA
        for i, (src, dst, note) in enumerate(gop, 4):
            for j, v in enumerate((src, dst, note), 4):
                cfg_ws.cell(i, j, v).font = F_DATA

        # --- Qua han 13111 (tong hop theo khach hang) ----------------------------
        step(80, "Tong hop theo khach hang")
        cols13 = [("Số tiền còn lại phải thu", "W", True), ("Tổng số quá hạn", "X", True)] + \
            [(b, c, True) for b, c in zip(BUCKETS, ["Y", "Z", "AA", "AB", "AC", "AD"])]
        con13 = dai_ly_theo((clean_name(r[RAW["ten_kh"]]), r[RAW["ten_daily"]])
                            for r in union_rows if can_tach(r[RAW["ten_kh"]]))
        ghi_qua_han(wb, "Quá hạn 13111",
                    f"TỔNG HỢP CÔNG NỢ THEO KHÁCH HÀNG – TK 13111 – đến {rdate:%d/%m/%Y}", "",
                    bc_name, bc_last, L_STD, L_DL, cols13, khs, con13)

    if co51:
        # --- Qua han 0051: nhu tren nhung bo so don goc co tong cong no = 0 --------
        step(84, "Tong hop qua han 0051")
        cols51 = [("Số tiền còn lại phải thu", "W", True), ("Tổng số quá hạn", "X", False)] + \
            [(b, c, False) for b, c in zip(BUCKETS, ["Z", "AA", "AB", "AC", "AD", "AE"])]
        khs51 = {}
        for r in rows51_no:
            k = clean_name(r[RAW["ten_kh"]])
            khs51.setdefault(k.upper(), k)
        con51 = dai_ly_theo((clean_name(r[RAW["ten_kh"]]), r[RAW["ten_daily"]])
                            for r in rows51_no if can_tach(r[RAW["ten_kh"]]))
        ghi_qua_han(wb, "Quá hạn 0051",
                    f"TỔNG HỢP CÔNG NỢ THEO KHÁCH HÀNG, TK 0051, đến {rdate:%d/%m/%Y}",
                    f"Không tính số đơn BH gốc có tổng công nợ bằng 0 (sheet 0051, cột TÌNH TRẠNG TT"
                    f" = \"{TT_KHONG_NO}\")",
                    "0051", last51, L(c51_kh), L(c51_dl), cols51,
                    sorted(khs51.values(), key=vn_key), con51,
                    them=f",'0051'!${L(cTT)}${ROW_DATA}:${L(cTT)}${last51},\"<>{TT_KHONG_NO}\"")

    if co13:
        # --- Theo ma KH: 3 khoi theo nhom nghiep vu ------------------------------
        step(88, "Dung 3 khoi 'Theo ma KH'")
        tk = wb.create_sheet("Theo mã KH")
        tk.sheet_view.showGridLines = False
        for j, w in enumerate([52, 18, 18, 46, 22], 1):
            tk.column_dimensions[L(j)].width = w
        groups = GROUPS
        by_group = {g: {} for g in groups}
        cap_tk = []
        for r in union_rows:
            g = resolve_group(r[RAW["ten_donbh"]])
            kh, note = resolve_customer(r[RAW["ten_kh"]])
            by_group[g].setdefault(kh, note)
            if can_tach(kh):
                cap_tk.append(((g, kh), r[RAW["ten_daily"]]))
        con_tk = dai_ly_theo(cap_tk)
        f10, f10i = Font(name=TNR, size=10), Font(name=TNR, size=10, italic=True)

        def rng_bc(c):
            return f"'{bc_name}'!${c}${ROW_DATA}:${c}${bc_last}"

        row, block_totals = 1, []
        for g in groups:
            tk.cell(row, 1, f"CÔNG NỢ QUÁ HẠN {g.upper()}").font = F_TITLE
            hrow = row + 2
            for j, hdr in enumerate(["Khách hàng", "Số tiền (VND)", "Tổng số quá hạn",
                                     "Giải trình", "Ghi chú"], 1):
                c = tk.cell(hrow, j, hdr)
                c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
                c.alignment = Alignment(horizontal="center", vertical="center")
            i, dong_con = hrow + 1, []
            for kh in sorted(by_group[g], key=vn_key):
                cha = i
                tk.cell(i, 1, kh).font = f10
                for j, sc in ((2, "W"), (3, "X")):
                    c = tk.cell(i, j, f"=SUMIFS({rng_bc(sc)},{rng_bc(L_KH)},$A{i},"
                                      f"{rng_bc(L_GRP)},\"{g}\")")
                    c.font, c.number_format = f10, NUMFMT["money"]
                tk.cell(i, 4).fill = FILL_INPUT
                tk.cell(i, 5, by_group[g][kh]).font = f10
                i += 1
                for dl in con_tk.get((g, kh), []):
                    ghi_dong_con(tk, i, dl, f10i)
                    for j, sc in ((2, "W"), (3, "X")):
                        c = tk.cell(i, j, f"=SUMIFS({rng_bc(sc)},{rng_bc(L_KH)},$A${cha},"
                                          f"{rng_bc(L_GRP)},\"{g}\",{rng_bc(L_DL)},{tieu_chi(dl)})")
                        c.font, c.number_format = f10i, NUMFMT["money"]
                    tk.cell(i, 4).fill = FILL_INPUT
                    dong_con.append(i)
                    i += 1
            trow = i
            tk.cell(trow, 1, "Tổng cộng").font = F_TOTAL
            for j in (2, 3):
                c = tk.cell(trow, j, f"=SUM({L(j)}{hrow + 1}:{L(j)}{trow - 1})"
                                     f"{tru_dong_con(L(j), dong_con)}")
                c.font, c.number_format, c.border = F_TOTAL, NUMFMT["money"], BOX
            block_totals.append(trow)
            row = trow + 5

        tk.cell(row, 1, "KIỂM TRA: tổng 3 khối − tổng BC 13111 (phải = 0)").font = F_TOTAL
        for j, sc in ((2, "W"), (3, "X")):
            c = tk.cell(row, j, "=" + "+".join(f"{L(j)}{t}" for t in block_totals) +
                        f"-'{bc_name}'!${sc}${ROW_TOTAL}")
            c.font, c.number_format, c.border = F_TOTAL, NUMFMT["money"], BOX

    step(95, "Ghi file")
    order = ["12", "13", "14", "15", MERGE_SHEET, "Quá hạn 13111", "Theo mã KH", "Quá hạn 0051",
             "0051", "Config"]
    wb._sheets = [wb[s] for s in order if s in wb.sheetnames]
    out = out_dir / f"{TEN_BC[phan]}-{rdate:%d.%m.%Y}.xlsx"
    wb.save(out)
    log(f"\nDa ghi: {out}")

    if co13:
        # --- canh bao gop khach hang con bo sot ----------------------------------
        parents = {norm(k) for k in khs}
        sot = []
        for k in khs:
            m = re.match(r"^\s*Khách lẻ.*?\((.+)\)\s*$", k, re.IGNORECASE | re.DOTALL)
            if m and norm(m.group(1)) in parents and norm(resolve_customer(k)[0]) == norm(k):
                sot.append((k, m.group(1).strip()))
        if sot:
            log("\n[CAN XEM] dong 'Khach le ... (X)' ma X cung la mot khach hang rieng, muon gop"
                f" thi them 1 dong vao sheet 'Gop khach hang' trong {CONFIG_NAME}:")
            for k, p in sot:
                log(f"    {k[:60]}  ->  {p[:45]}")
        wild = [k for k in khs if re.search(r"[*?~]", k)]
        if wild:
            log("\n[CANH BAO] ten khach hang co ky tu dai dien (* ? ~), SUMIFS co the cong nham:")
            for k in wild:
                log(f"    {k[:70]}")
    return out


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else app_dir() / DIR_IN
    build(src, sys.argv[2] if len(sys.argv) > 2 else None)
