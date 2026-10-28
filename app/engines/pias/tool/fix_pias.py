"""PIAS fix - PVI khac.

Nhan dup Run_PIAS_Fix.bat -> chon THU MUC input chua file thu PIAS (ctXXXX.xlsx) va
file danh sach PVI khac cua chi Kieu; tool tu nhan loai file theo tieu de. Keo tha
thu muc (hoac cac file) vao .bat thi bo qua hop thoai. Ket qua: output\\<ngay gio>.

    py -3.13 tool\\fix_pias.py [thu_muc | pias.xlsx ... kieu.xlsx] [--sheet TEN] [--date 08/09,09/09]

Dong PIAS khop (dien giai bo so thu tu + so tien) voi mot mon trong file chi Kieu:
Tk co 33889 -> 33664, Khach co + P/Ban co -> ma don vi, xoa Can bo co, to vang o Tk co.
Mon chi Kieu tach tien thanh nhieu dong thi cong theo (ngay, noi dung) roi moi so.
"""
import argparse
import datetime as dt
import io
import os
import re
import sys
import time
import unicodedata
import warnings
from collections import Counter, defaultdict
from functools import reduce
from itertools import product
from math import gcd

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import openpyxl                                              # noqa: E402
from openpyxl.styles import PatternFill                      # noqa: E402
from openpyxl.utils import get_column_letter                 # noqa: E402

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
UNIT_FILE = os.path.join(ROOT, "Unit_Codes.xlsx")
OUT_ROOT = os.path.join(ROOT, "output")
LAST_DIR = os.path.join(HERE, "last_folder.txt")

TK_OLD, TK_NEW = "33889", "33664"
DIGITAL = "00.18450000"
YELLOW = PatternFill("solid", start_color="FFFFFF00")
CODE_RE = re.compile(r"^\d{2}\.\d{8}$")
BAR_WIDTH = 30
LINE_W = 96
SHOW = 15          # so dong canh bao in ra man hinh moi nhom, con lai xem log.txt


class ToolError(Exception):
    pass


# ------------------------------------------------------------ thanh tien trinh

def _enable_ansi():
    if not sys.stdout.isatty():
        return False
    try:
        import ctypes
        k = ctypes.windll.kernel32
        h = k.GetStdHandle(-11)
        mode = ctypes.c_uint32()
        if k.GetConsoleMode(h, ctypes.byref(mode)):
            k.SetConsoleMode(h, mode.value | 0x0004)
            return True
    except Exception:
        pass
    return False


COLOR = _enable_ansi()
C_ON = "\033[96m" if COLOR else ""
C_OK = "\033[92m" if COLOR else ""
C_WARN = "\033[93m" if COLOR else ""
C_OFF = "\033[0m" if COLOR else ""


class Progress:
    """Thanh chay tai cho: [######....]  3/9  <viec dang lam>. step() goi TRUOC
    moi viec de nguoi dung thay tool dang lam gi (doc file chi Kieu mat vai giay)."""

    def __init__(self, total):
        self.total = max(1, total)
        self.n = 0
        self.t0 = time.time()
        self.live = sys.stdout.isatty()

    def step(self, label=""):
        self.n += 1
        if not self.live:
            return
        fill = int(round(BAR_WIDTH * min(self.n, self.total) / self.total))
        bar = "█" * fill + "░" * (BAR_WIDTH - fill)
        plain = "   %s  %*d/%d  %s" % (bar, len(str(self.total)), self.n, self.total, label)
        if len(plain) > LINE_W:
            plain = plain[:LINE_W - 3] + "..."
        sys.stdout.write("\r" + plain.replace(bar, C_ON + bar + C_OFF, 1) + " " * (LINE_W - len(plain)))
        sys.stdout.flush()

    def wipe(self):
        if self.live:
            sys.stdout.write("\r" + " " * LINE_W + "\r")
            sys.stdout.flush()

    def finish(self, msg):
        if self.live:
            bar = "█" * BAR_WIDTH
            plain = "   %s  %d/%d" % (bar, self.total, self.total)
            sys.stdout.write("\r" + plain.replace(bar, C_ON + bar + C_OFF, 1)
                             + " " * max(0, LINE_W - len(plain)) + "\n")
            sys.stdout.flush()
        print("   %s>%s %s  (%.1fs)" % (C_OK, C_OFF, msg, time.time() - self.t0))


# --------------------------------------------------------- hop thoai chon thu muc

DIALOG_TITLE = "Select the input folder (PIAS file + PVI khac list)"


def _remember(folder):
    try:
        with io.open(LAST_DIR, "w", encoding="utf-8") as f:
            f.write(folder)
    except OSError:
        pass


def _start_dir():
    """Thu muc mo san trong hop thoai: lan truoc, roi thu muc input, roi thu muc goc."""
    try:
        with io.open(LAST_DIR, encoding="utf-8") as f:
            d = f.read().strip()
        if d and os.path.isdir(d):
            return d
    except OSError:
        pass
    d = os.path.join(ROOT, "input")
    return d if os.path.isdir(d) else ROOT


def _ask_com():
    """Hop thoai chon thu muc kieu Explorer (IFileOpenDialog + FOS_PICKFOLDERS) goi
    thang qua COM bang ctypes: ban Python nhung khong co tkinter."""
    import ctypes
    from ctypes import POINTER, byref, c_long, c_ulong, c_void_p, c_wchar_p, wintypes

    ole32, shell32 = ctypes.windll.ole32, ctypes.windll.shell32

    class GUID(ctypes.Structure):
        _fields_ = [("a", wintypes.DWORD), ("b", wintypes.WORD), ("c", wintypes.WORD),
                    ("d", ctypes.c_ubyte * 8)]

    def guid(s):
        g = GUID()
        ole32.CLSIDFromString(c_wchar_p(s), byref(g))
        return g

    def call(obj, idx, *args):
        """Goi ham thu idx trong vtable; args la cap (kieu, gia tri). Tra HRESULT."""
        vtbl = ctypes.cast(obj, POINTER(POINTER(c_void_p))).contents
        fn = ctypes.WINFUNCTYPE(c_long, c_void_p, *[t for t, _ in args])(vtbl[idx])
        return fn(obj, *[v for _, v in args]) & 0xFFFFFFFF

    def release(obj):
        vtbl = ctypes.cast(obj, POINTER(POINTER(c_void_p))).contents
        ctypes.WINFUNCTYPE(c_ulong, c_void_p)(vtbl[2])(obj)

    ole32.CoInitialize(None)
    dlg = c_void_p()
    hr = ole32.CoCreateInstance(byref(guid("{DC1C5A9C-E88A-4dde-A5A1-60F82A20AEF7}")), None, 1,
                                byref(guid("{d57c7288-d4ad-4768-be02-9d969532d960}")), byref(dlg))
    if hr != 0 or not dlg:
        raise OSError("CoCreateInstance FileOpenDialog failed")
    try:
        opts = wintypes.DWORD()
        call(dlg, 10, (POINTER(wintypes.DWORD), byref(opts)))                      # GetOptions
        call(dlg, 9, (wintypes.DWORD, opts.value | 0x20 | 0x40 | 0x800))          # PICKFOLDERS|FORCEFILESYSTEM|PATHMUSTEXIST
        call(dlg, 17, (c_wchar_p, DIALOG_TITLE))                                   # SetTitle
        start = c_void_p()
        if shell32.SHCreateItemFromParsingName(c_wchar_p(_start_dir()), None,
                                               byref(guid("{43826d1e-e718-42ee-bc55-a1e261c37bfe}")),
                                               byref(start)) == 0 and start:
            call(dlg, 12, (c_void_p, start))                                       # SetFolder
            release(start)
        hr = call(dlg, 3, (wintypes.HWND, None))                                   # Show
        if hr == 0x800704C7:                                                       # bam Cancel
            return ""
        if hr != 0:
            raise OSError("IFileOpenDialog.Show 0x%08X" % hr)
        item = c_void_p()
        if call(dlg, 20, (POINTER(c_void_p), byref(item))) != 0:                   # GetResult
            raise OSError("IFileOpenDialog.GetResult failed")
        name = c_wchar_p()
        hr = call(item, 5, (wintypes.DWORD, 0x80058000), (POINTER(c_wchar_p), byref(name)))  # GetDisplayName(FILESYSPATH)
        release(item)
        if hr != 0:
            raise OSError("IShellItem.GetDisplayName failed")
        path = name.value
        ole32.CoTaskMemFree(name)
        return path
    finally:
        release(dlg)


def _ask_tk():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    d = filedialog.askdirectory(parent=root, title=DIALOG_TITLE, initialdir=_start_dir())
    root.destroy()
    return d or ""


def ask_folder():
    """Duong dan thu muc, "" neu bam Cancel. Khong mo duoc hop thoai nao thi hoi o cua so lenh."""
    for fn in (_ask_com, _ask_tk):
        try:
            return fn()
        except Exception:
            continue
    try:
        s = input("Paste the input folder path (Enter = %s): " % _start_dir()).strip().strip('"')
    except EOFError:
        s = ""
    return s or _start_dir()


def check_path(p):
    p = os.path.abspath(p)
    if not os.path.isfile(p):
        raise ToolError("File not found: %s" % p)
    if not p.lower().endswith((".xlsx", ".xlsm")):
        raise ToolError("%s is not .xlsx - open it in Excel and Save As .xlsx first." % os.path.basename(p))
    return p


def excel_in(folder):
    """File Excel trong thu muc (khong vao thu muc con), bo file tam ~$ va file _Fixed da xuat."""
    out = []
    for n in sorted(os.listdir(folder)):
        stem, ext = os.path.splitext(n)
        if ext.lower() in (".xlsx", ".xlsm") and not n.startswith("~$") \
                and not stem.lower().endswith("_fixed") and os.path.isfile(os.path.join(folder, n)):
            out.append(os.path.join(folder, n))
    return out


def pick_inputs(args):
    """Tra ve (file PIAS, danh sach file chi Kieu). Khong truyen gi thi hoi thu muc;
    keo tha thu muc hay file vao .bat deu nhan. Loai file nhan theo tieu de."""
    if not args:
        folder = ask_folder()
        if not folder:
            raise ToolError("Cancelled - no folder selected.")
        args = [folder]
    paths = []
    for a in args:
        p = os.path.abspath(a)
        if os.path.isdir(p):
            found = excel_in(p)
            if not found:
                raise ToolError("No Excel file in %s" % p)
        else:
            found = [check_path(p)]
        paths += [f for f in found if f not in paths]
    where = os.path.dirname(paths[0])
    pias = [p for p in paths if is_pias(p)]
    lists = [p for p in paths if p not in pias]
    if not pias:
        raise ToolError("No PIAS file (columns 'Diễn giải' and 'Tk có') in %s" % where)
    if not lists:
        raise ToolError("No PVI khac list from chị Kiều in %s" % where)
    _remember(os.path.abspath(args[0]) if os.path.isdir(args[0]) else where)
    return pias, lists


# --------------------------------------------------------------- tien ich nho

def key(s):
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    s = re.sub(r"\btp\b\.?", "thanh pho", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


MOJIBAKE = re.compile(r"[\xc2-\xdf][\x80-\xbf]|[\xe0-\xef][\x80-\xbf]{2}")


def _unmojibake(m):
    try:
        return m.group(0).encode("latin-1").decode("utf-8")
    except UnicodeError:
        return m.group(0)


def norm(s):
    """Noi dung de so: PIAS xuat hong dau ("700Ä\\x90" = "700Đ") va danh sach doi dau cham
    phay ("STB;39598888;" vs "STB,39598888,", CT0820 18/09) nen chi so chu + so, bo dau."""
    s = MOJIBAKE.sub(_unmojibake, str(s or ""))
    return key(s).upper()


def same_content(p, k):
    if p == k:
        return True
    if min(len(p), len(k)) < 12:
        return False
    if k.startswith(p) or p.startswith(k):
        return True
    # lech vai ky tu cuoi (ky tu hong PIAS khong sua duoc): chung phan dau gan het ban ngan
    n = len(os.path.commonprefix([p, k]))
    return n >= 40 and n >= min(len(p), len(k)) - 2


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.fullmatch(r"\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*", str(v or ""))
    if m:
        d, mo, y = map(int, m.groups())
        try:
            return dt.date(y + 2000 if y < 100 else y, mo, d)
        except ValueError:
            return None
    return None


def to_amount(v):
    if isinstance(v, (bool, dt.datetime, dt.date)):
        return None
    if isinstance(v, (int, float)):
        return round(v)
    s = str(v or "").replace("\xa0", " ").strip()
    if re.fullmatch(r"-?\d{1,3}([.,\s]\d{3})+|-?\d+", s):
        return int(re.sub(r"\D", "", s)) * (-1 if s.startswith("-") else 1)
    return None


def money(n):
    return f"{n:,.0f}".replace(",", " ")


def ddmm(d):
    return d.strftime("%d/%m") if d else "--"


def load_units():
    if not os.path.isfile(UNIT_FILE):
        raise ToolError("Unit_Codes.xlsx is missing next to Run_PIAS_Fix.bat.")
    ws = openpyxl.load_workbook(UNIT_FILE, read_only=True).worksheets[0]
    units = {}
    for name, code in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        if not name or not code:
            continue
        k = key(name)
        units.setdefault(k, str(code).strip())
        short = re.sub(r"^(cong ty |bao hiem |pvi )+", "", k)
        if short:
            units.setdefault(short, str(code).strip())
    return units


def unit_code(v, units):
    s = str(v or "").strip()
    if CODE_RE.match(s):
        return s
    k = key(s)
    return units.get(k) or units.get(re.sub(r"^(cong ty |bao hiem |pvi )+", "", k))


# ---------------------------------------------------------------------- PIAS

def is_pias(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.worksheets[0]
        head = {key(v) for row in ws.iter_rows(max_row=5, max_col=15, values_only=True) for v in row}
        wb.close()
        return {"dien giai", "tk co"} <= head
    except Exception:
        return False


def read_pias(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    want = {"text": r"dien giai", "tk": r"tk co$", "amt": r"so tien", "khach": r"khach co",
            "canbo": r"can bo co", "pban": r"p ban co"}
    for hr in range(1, 6):
        cols = {}
        for c in ws[hr]:
            for k, pat in want.items():
                if k not in cols and re.match(pat, key(c.value)):
                    cols[k] = c.column
        if {"text", "amt"} <= cols.keys():
            break
    else:
        raise ToolError("%s: header row with 'Diễn giải' / 'Số tiền' not found." % os.path.basename(path))
    missing = want.keys() - cols.keys()
    if missing:
        raise ToolError("%s: missing column(s) %s" % (os.path.basename(path), ", ".join(sorted(missing))))
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        txt = ws.cell(r, cols["text"]).value
        amt = to_amount(ws.cell(r, cols["amt"]).value)
        if txt is None or amt is None:
            continue
        m = re.match(r"\s*(\d+)\s*\.\s*", str(txt))
        rows.append({"r": r, "stt": m.group(1) if m else "row %d" % r, "amt": amt,
                     "c": norm(str(txt)[m.end():] if m else txt)})
    # CT0791 (17/09/2026) xuat tu PIAS bi hong cot tien: 239/256 dong cung 26 790 000.
    # Mot ngay binh thuong so tien hay gap nhat chi ~12% so dong.
    if len(rows) >= 20:
        amt, n = Counter(p["amt"] for p in rows).most_common(1)[0]
        if n > len(rows) / 2:
            raise ToolError("%s: %d of %d rows have the same amount %s in column 'Số tiền VNĐ' - the file "
                            "looks broken, do not upload it. Export it from PIAS again and check that column."
                            % (os.path.basename(path), n, len(rows), money(amt)))
    return wb, ws, cols, rows


# -------------------------------------------------------- file chi Kieu

HINT = {"date": r"ngay", "content": r"noi dung|dien giai|sao ke", "amount": r"tien", "unit": r"don vi"}


def sheet_candidates(ws, units):
    """Dau vao khong co dinh: tim dong tieu de, phan loai cot theo DU LIEU
    (ngay / so / chu dai / ten don vi), roi thu hep theo tu khoa tieu de."""
    data = [list(r) for r in ws.iter_rows(max_col=80, values_only=True)]
    if len(data) < 2:
        return None
    hint_all = "|".join(HINT.values())

    def head_score(row):
        strs = [key(v) for v in row if isinstance(v, str) and len(v) <= 60]
        return sum(10 for s in strs if re.search(hint_all, s)) + len(strs)

    hr = max(range(min(20, len(data))), key=lambda i: head_score(data[i]))
    header = [key(v) if isinstance(v, str) else "" for v in data[hr]]
    body = [r for r in data[hr + 1:] if sum(v is not None for v in r) >= 2]
    if not body:
        return None
    ncol = max(len(r) for r in body)
    kinds = defaultdict(list)
    unit_ratio = {}
    for c in range(ncol):
        vals = [r[c] for r in body if c < len(r) and r[c] not in (None, "")]
        if len(vals) < max(3, len(body) * 0.3):
            continue
        n = len(vals)
        if sum(to_date(v) is not None for v in vals) >= 0.6 * n:
            kinds["date"].append(c)
            continue
        if sum(to_amount(v) is not None for v in vals) >= 0.6 * n:
            kinds["amount"].append(c)
        strs = [v for v in vals if isinstance(v, str)]
        if len(strs) >= 0.6 * n and sum(map(len, strs)) / len(strs) >= 15:
            kinds["content"].append(c)
        unit_ratio[c] = sum(unit_code(v, units) is not None for v in vals) / n
    for k in kinds:
        hinted = [c for c in kinds[k] if c < len(header) and re.search(HINT[k], header[c])]
        if hinted:
            kinds[k] = hinted
    unit_cols = [c for c, x in unit_ratio.items() if x >= 0.6]
    if not unit_cols or not kinds["content"] or not kinds["amount"]:
        return None
    unit_col = max(unit_cols, key=lambda c: (unit_ratio[c], bool(re.search(HINT["unit"], header[c] if c < len(header) else ""))))
    codes = [unit_code(r[unit_col], units) for r in body if unit_col < len(r) and r[unit_col]]
    digital = sum(c == DIGITAL for c in codes) / max(1, sum(c is not None for c in codes))
    return {"hr": hr, "header": list(data[hr]), "body": body, "date": kinds["date"] or [None],
            "content": kinds["content"], "amount": kinds["amount"], "unit": unit_col, "digital": digital}


def build_groups(sc, dc, cc, ac, units, only_dates=None):
    groups = {}
    for r in sc["body"]:
        def get(c):
            return r[c] if c is not None and c < len(r) else None
        amt, txt = to_amount(get(ac)), get(cc)
        d = to_date(get(dc)) if dc is not None else None
        if amt is None or not txt or (dc is not None and d is None):
            continue
        if only_dates and (d is None or (d.day, d.month) not in only_dates):
            continue
        u = get(sc["unit"])
        if unit_code(u, units) == DIGITAL:
            continue
        g = groups.setdefault((d, norm(txt)), {"d": d, "c": norm(txt), "raw": str(txt), "amt": 0, "n": 0,
                                                "lines": [], "codes": set(), "names": set()})
        g["amt"] += amt
        g["n"] += 1
        g["lines"].append(amt)
        g["names"].add(str(u or "").strip())
        g["codes"].add(unit_code(u, units))
    return groups


def repeats(g):
    """So lan bo dong cua mot mon lap lai y het (ucln cua so lan xuat hien moi so tien)."""
    return reduce(gcd, Counter(g["lines"]).values()) if g["lines"] else 0


def match(pias, groups):
    """Moi mon chi Kieu gan cho dung 1 dong PIAS; khop nhieu thi de lai bao sua tay.
    Tra ve (gan, [(dong chua chac, cac mon ung vien)], [nhom dong PIAS cung noi dung sua chung])."""
    by_amt = defaultdict(list)
    for g in groups.values():
        by_amt[g["amt"]].append(g)
    cand = {p["r"]: [g for g in by_amt[p["amt"]] if same_content(p["c"], g["c"])] for p in pias}
    claims = Counter(id(g) for gs in cand.values() for g in gs)
    assign, used = {}, set()
    for p in pias:
        gs = cand[p["r"]]
        if len(gs) == 1 and claims[id(gs[0])] == 1:
            assign[p["r"]] = gs[0]
            used.add(id(gs[0]))
    top = Counter(g["d"] for g in assign.values())
    for p in pias:
        if p["r"] in assign:
            continue
        gs = [g for g in cand[p["r"]] if id(g) not in used]
        if len(gs) > 1 and top:
            best = max(top[g["d"]] for g in gs)
            gs = [g for g in gs if top[g["d"]] == best]
        if len(gs) == 1 and sum(id(gs[0]) in {id(x) for x in cand[q["r"]]} for q in pias if q["r"] not in assign) == 1:
            assign[p["r"]] = gs[0]
            used.add(id(gs[0]))
    # PIAS co nhieu dong cung noi dung: cong lai roi so tong
    rest = defaultdict(list)
    for p in pias:
        if p["r"] not in assign:
            rest[p["c"]].append(p)
    together = []
    for c, ps in rest.items():
        if len(ps) < 2:
            continue
        total = sum(p["amt"] for p in ps)
        gs = [g for g in by_amt[total] if id(g) not in used and same_content(c, g["c"])]
        if len(gs) == 1:
            for p in ps:
                assign[p["r"]] = gs[0]
            used.add(id(gs[0]))
            together.append((ps, gs[0], "pias"))
    # Danh sach nhap trung: cac dong cua mon lap lai dung m lan (11 830 000 x2, hay ca cap
    # 480 700 + 50 000 x2), PIAS chi co k < m dong bang 1 ban cung noi dung (tong khong
    # khop nen cac buoc tren bo qua). Dong PIAS khop dung 1 ban nen van sua, va bao trung.
    for g in groups.values():
        m = repeats(g)
        if id(g) in used or m < 2:
            continue
        ps = [p for p in pias if p["r"] not in assign and not cand[p["r"]]
              and p["amt"] == g["amt"] // m and same_content(p["c"], g["c"])]
        if 0 < len(ps) < m:
            for p in ps:
                assign[p["r"]] = g
            used.add(id(g))
            together.append((ps, g, "list"))
    # con ung vien ma khong gan duoc = khop nhieu mon, hoac nhieu dong PIAS tranh 1 mon (dong trung)
    unsure = [(p, cand[p["r"]]) for p in pias if p["r"] not in assign and cand[p["r"]]]
    return assign, unsure, together


def pick_kieu(path, pias_all, units, sheet=None, only_dates=None):
    """Thu moi to hop cot, lay to hop khop duoc nhieu dong PIAS nhat."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible" or (sheet and ws.title != sheet):
            continue
        sc = sheet_candidates(ws, units)
        if not sc:
            continue
        for dc, cc, ac in product(sc["date"], sc["content"], sc["amount"]):
            if cc == ac:
                continue
            groups = build_groups(sc, dc, cc, ac, units, only_dates)
            ok = sum(None not in g["codes"] and len(g["codes"]) == 1 for g in match(pias_all, groups)[0].values())
            # file chi Kieu la danh sach PVI khac: sheet chi tiet don (da phan PVI Digital)
            # khop nhieu dong hon nhung chi dung khi khong con sheet nao khac
            score = (sc["digital"] <= 0.5, ok)
            if ok and (not best or score > best[0]):
                best = (score, ws.title, sc, dc, cc, ac)
    wb.close()
    return best


def diagnose(path, pias_all, units, sheet=None, only_dates=None):
    """Chi chay khi khong khop duoc dong nao: noi ro vi sao (khong mo duoc file / khong
    do ra cot / danh sach khong co ngay cua file PIAS / trung noi dung nhung lech tien)."""
    name = os.path.basename(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        msg = "%s: %s" % (e.__class__.__name__, str(e)[:100])
        hint = ""
        if "zip" in msg.lower() or "invalidfile" in msg.lower():
            hint = ("\n      The file looks protected (password, or a sensitivity label with encryption)."
                    "\n      Copy its sheet into a new Excel workbook, save as .xlsx without protection, run again.")
        return ["%s: cannot open it (%s)%s" % (name, msg, hint)]
    out, seen = [], []
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible" or (sheet and ws.title != sheet):
                continue
            seen.append(ws.title)
            sc = sheet_candidates(ws, units)
            if not sc:
                continue
            dc, cc, ac = sc["date"][0], sc["content"][0], sc["amount"][0]
            groups = build_groups(sc, dc, cc, ac, units, only_dates)
            days = sorted(Counter(g["d"] for g in groups.values()).items(), key=lambda x: x[0] or dt.date.min)
            out.append("%s, sheet '%s': date %s | content %s | amount %s | unit %s"
                       % (name, ws.title, col_name(sc, dc), col_name(sc, cc), col_name(sc, ac), col_name(sc, sc["unit"])))
            if not groups:
                out.append("   no item of another PVI unit in this sheet")
                continue
            out.append("   days in the list: %s .. %s, last days: %s"
                       % (ddmm(days[0][0]), ddmm(days[-1][0]),
                          ", ".join("%s (%d)" % (ddmm(d), n) for d, n in days[-5:])))
            loose = [(p, g) for p in pias_all for g in groups.values() if same_content(p["c"], g["c"])]
            if loose:
                out.append("   %d PIAS row(s) have the same content as an item but a different amount, e.g.:" % len(loose))
                for p, g in loose[:5]:
                    out.append("      STT %s: PIAS %s vs list %s (%s, %s)"
                               % (p["stt"], money(p["amt"]), money(g["amt"]), ddmm(g["d"]), g["raw"][:40]))
            else:
                out.append("   no PIAS row has the same content as any item: the PIAS file is probably for a day "
                           "the list does not have yet")
    finally:
        wb.close()
    if not out:
        out.append("%s: no sheet with the 4 columns payment date, statement content, amount, PVI unit name "
                   "(visible sheets: %s)" % (name, ", ".join(seen) or "none"))
    return out


def col_name(sc, c):
    if c is None:
        return "(none)"
    h = sc["header"][c] if c < len(sc["header"]) else None
    return get_column_letter(c + 1) + (" '%s'" % h if isinstance(h, str) else "")


def out_path(run_dir, src, taken):
    m = re.match(r"[A-Za-z]+\d+", os.path.basename(src))
    base = "%s_Fixed" % (m.group(0) if m else "PIAS")
    name, i = base, 2
    while name in taken:
        name, i = "%s_%d" % (base, i), i + 1
    taken.add(name)
    return os.path.join(run_dir, name + ".xlsx")


# --------------------------------------------------------------------------- chay

def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="*", help="input folder, or the PIAS file(s) + the PVI khac list")
    ap.add_argument("--sheet", help="sheet of the PVI khac list, when the guess is wrong")
    ap.add_argument("--date", help="only these days of the PVI khac list, e.g. 08/09,09/09")
    a = ap.parse_args(argv)

    print("=== PIAS fix - PVI khac ===")
    pias_paths, lists = pick_inputs(a.inputs)
    print("Folder   : %s" % os.path.dirname(pias_paths[0]))
    for i, p in enumerate(pias_paths):
        print("%s%s" % ("PIAS     : " if i == 0 else "           ", os.path.basename(p)))
    print("")

    only_dates = None
    if a.date:
        try:
            only_dates = {(int(d), int(m)) for d, m in (x.strip().split("/")[:2] for x in a.date.split(","))}
        except ValueError:
            raise ToolError("Cannot read --date %s (use 08/09,09/09)" % a.date)

    bar = Progress(len(pias_paths) * 2 + len(lists) + 2)
    bar.step("reading Unit_Codes.xlsx")
    units = load_units()
    files = []
    for p in pias_paths:
        bar.step("reading %s" % os.path.basename(p))
        files.append((p,) + read_pias(p))
    pias_all = [row for f in files for row in f[4]]

    # Nhieu ban danh sach (chi Kieu gui ban cap nhat): lay ban khop nhieu dong nhat,
    # bang nhau thi lay ban sua gan nhat.
    best, kieu, unused = None, None, []
    for p in lists:
        bar.step("reading %s" % os.path.basename(p))
        try:
            got = pick_kieu(p, pias_all, units, a.sheet, only_dates)
        except Exception:
            got = None
        if got and (not best or (got[0], os.path.getmtime(p)) > (best[0], os.path.getmtime(kieu))):
            if kieu:
                unused.append(kieu)
            best, kieu = got, p
        else:
            unused.append(p)
    if not best:
        bar.wipe()
        print("")
        print("%s!! No PIAS row matches the PVI khac list. What the tool found:%s" % (C_WARN, C_OFF))
        print("   PIAS: %s" % ", ".join("%s (%d rows)" % (os.path.basename(f[0]), len(f[4])) for f in files))
        for p in lists:
            for s in diagnose(p, pias_all, units, a.sheet, only_dates):
                print("   " + s)
        raise ToolError("nothing fixed, see above.")
    _, title, sc, dc, cc, ac = best
    groups = build_groups(sc, dc, cc, ac, units, only_dates)

    now = dt.datetime.now()
    base_dir = os.path.join(OUT_ROOT, now.strftime("%Y-%m-%d %H%M"))
    run_dir, i = base_dir, 1
    while os.path.isdir(run_dir):
        i += 1
        run_dir = "%s (%d)" % (base_dir, i)
    os.makedirs(run_dir)

    taken, used_all, flagged, dates_hit = set(), set(), set(), Counter()
    results, log = [], []
    for path, wb, ws, cols, rows in files:
        bar.step("matching %s" % os.path.basename(path))
        assign, unsure, together = match(rows, groups)
        fixed, done, check, notes = [], [], [], []
        for p in rows:
            g = assign.get(p["r"])
            if not g:
                continue
            used_all.add(id(g))
            dates_hit[g["d"]] += 1
            codes = g["codes"] - {None}
            names = ", ".join(sorted(g["names"]))
            if None in g["codes"] or len(codes) != 1:
                check.append("STT %s (%s): unit '%s' has no code - add it to Unit_Codes.xlsx and run again"
                             % (p["stt"], money(p["amt"]), names))
                continue
            code = codes.pop()

            def cell(k):
                return ws.cell(p["r"], cols[k])
            tk = str(cell("tk").value or "").strip()
            if tk == TK_NEW and str(cell("khach").value or "").strip() == code \
                    and str(cell("pban").value or "").strip() == code:
                done.append(p)
                continue
            if tk not in (TK_OLD, TK_NEW):
                check.append("STT %s (%s): Tk có is %s, not %s - left as is" % (p["stt"], money(p["amt"]), tk, TK_OLD))
                continue
            cell("tk").value = TK_NEW
            cell("khach").value = code
            cell("canbo").value = None
            cell("pban").value = code
            cell("tk").fill = YELLOW            # nhu file chi sua tay: chi to o Tk co

            fixed.append((p, g, code))
        # dong trung: nhieu dong PIAS cung noi dung + so tien ma danh sach chi co 1 mon
        by_item = defaultdict(list)
        for p, gs in unsure:
            flagged.update(id(g) for g in gs)
            if len(gs) == 1:
                by_item[id(gs[0])].append((p, gs[0]))
            else:
                check.append("STT %s (%s): matches %d items of the PVI khac list (%s) - fix by hand"
                             % (p["stt"], money(p["amt"]), len(gs), ", ".join(ddmm(g["d"]) for g in gs)))
        for ps in by_item.values():
            g = ps[0][1]
            stts = " and ".join("STT %s" % p["stt"] for p, _ in ps)
            if len(ps) > 1:
                check.append("DUPLICATE %s (%s each): identical rows in PIAS, but the PVI khac list has "
                             "this item only once (%s) - none fixed, fix by hand" % (stts, money(g["amt"]), ddmm(g["d"])))
            else:
                check.append("%s (%s): its item in the PVI khac list also matches another row - fix by hand"
                             % (stts, money(g["amt"])))
        fixed_rows = {p["r"] for p, _, _ in fixed}
        for ps, g, kind in together:
            if not all(p["r"] in fixed_rows for p in ps):
                continue
            stts = " and ".join("STT %s" % p["stt"] for p in ps)
            if kind == "pias":
                notes.append("DUPLICATE %s: same content on %d PIAS rows, fixed together against %d line(s) "
                             "of the PVI khac list (total %s) - make sure it is not paid twice"
                             % (stts, len(ps), g["n"], money(g["amt"])))
            else:
                m = repeats(g)
                notes.append("DUPLICATE in the PVI khac list: %s (%s) fixed, but the list has this item %d times "
                             "(%d x %s, %s) while PIAS has %d row(s) - check the list"
                             % (stts, money(ps[0]["amt"]), m, m, money(g["amt"] // m), ddmm(g["d"]), len(ps)))

        out = None
        if fixed:
            out = out_path(run_dir, path, taken)
            try:
                wb.save(out)
            except PermissionError:
                raise ToolError("Cannot write %s - is it open in Excel?" % out)
        results.append((path, rows, fixed, done, check, notes, out))
        log.append("[%s]  %d rows | %d fixed | %d already fixed | %d to check%s"
                   % (os.path.basename(path), len(rows), len(fixed), len(done), len(check),
                      " | -> %s" % os.path.basename(out) if out else " | nothing to fix, no file written"))
        for p, g, code in fixed:
            log.append("   fixed  STT %5s | %14s | %s %s%s" % (p["stt"], money(p["amt"]), code, ", ".join(sorted(g["names"])),
                                                         " | %d lines in the list" % g["n"] if g["n"] > 1 else ""))
        if done:
            log.append("   already fixed: STT " + ", ".join(p["stt"] for p in done))
        log.extend("   CHECK  " + c for c in check)
        log.extend("   NOTE   " + c for c in notes)
        log.append("")

    days = {g["d"] for g in groups.values()} if only_dates else set(dates_hit)
    left = []
    for g in sorted((g for g in groups.values() if id(g) not in used_all | flagged and g["d"] in days),
                    key=lambda g: (g["d"] or dt.date.min, g["amt"])):
        near = ["STT %s amount %s" % (p["stt"], money(p["amt"])) for p in pias_all if same_content(p["c"], g["c"])]
        line = "%s | %12s | %-20s | %s" % (ddmm(g["d"]), money(g["amt"]), ", ".join(sorted(g["names"]))[:20], g["raw"][:60])
        m = repeats(g)
        if m > 1:
            line += "\n        DUPLICATE? the list has this item %d times (%d x %s)" % (m, m, money(g["amt"] // m))
        elif g["n"] > 1:
            line += "\n        the list has %d lines for it: %s" % (g["n"], " + ".join(money(x) for x in g["lines"]))
        if near:
            line += "\n        PIAS has the same content but a different amount: " + "; ".join(near)
        left.append(line)

    bar.step("writing log")
    head = [
        "PIAS     : " + "\n           ".join(pias_paths),
        "Run      : %s" % now.strftime("%d/%m/%Y %H:%M"),
        "PVI khac : %s" % os.path.basename(kieu)
        + "".join("\n           (not used: %s)" % os.path.basename(u) for u in unused),
        "List     : sheet '%s', header row %d" % (title, sc["hr"] + 1),
        "           date %s | content %s | amount %s | unit %s"
        % (col_name(sc, dc), col_name(sc, cc), col_name(sc, ac), col_name(sc, sc["unit"])),
        "Days     : " + ", ".join("%s (%d)" % (ddmm(d), n) for d, n in sorted(dates_hit.items(), key=lambda x: x[0] or dt.date.min)),
        "",
    ]
    with io.open(os.path.join(run_dir, "log.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(head + log))
        f.write("\nPVI khac items of these days NOT matched in the PIAS file (%d):\n" % len(left))
        f.write("\n".join(left) if left else "(none)")
        f.write("\n")
    bar.finish("%d file(s) written" % sum(1 for r in results if r[-1]))

    print("")
    print("\n".join(head[2:6]))
    for path, rows, fixed, done, check, notes, out in results:
        print("")
        print("[%s]  %d rows  |  %s%d fixed%s  |  %d already fixed"
              % (os.path.basename(path), len(rows), C_OK, len(fixed), C_OFF, len(done)))
        print("   -> %s" % (os.path.basename(out) if out else "nothing to fix, no file written"))
        if check:
            print("   %s!! %d to check by hand:%s" % (C_WARN, len(check), C_OFF))
            for c in check[:SHOW]:
                print("      - " + c)
            if len(check) > SHOW:
                print("      ... %d more, see log.txt" % (len(check) - SHOW))
        if notes:
            print("   %s!! %d fixed, but please look:%s" % (C_WARN, len(notes), C_OFF))
            for c in notes[:SHOW]:
                print("      - " + c)
            if len(notes) > SHOW:
                print("      ... %d more, see log.txt" % (len(notes) - SHOW))
    print("")
    if left:
        print("%s!! %d item(s) of the PVI khac list NOT matched in the PIAS file:%s" % (C_WARN, len(left), C_OFF))
        for s in left[:SHOW]:
            print("   - " + s)
        if len(left) > SHOW:
            print("   ... %d more, see log.txt" % (len(left) - SHOW))
    else:
        print("Every PVI khac item of these days is in the PIAS file.")
    print("")
    print("Output folder:")
    print("   %s" % run_dir)
    if not os.environ.get("PIASFIX_NO_OPEN"):
        try:
            os.startfile(run_dir)                            # noqa: S606
        except Exception:
            pass


if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except ToolError as e:
        if sys.stdout.isatty():
            sys.stdout.write("\r" + " " * LINE_W + "\r")
        print("")
        print("ERROR: %s" % e)
        sys.exit(1)
